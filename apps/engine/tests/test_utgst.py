from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

# GSTINs below are checksum-valid so gstin_flags() stays quiet and only the
# UTGST classification shows up in the Flags column.
CHANDIGARH_GSTIN = "04AABCU9603R1ZV"
KARNATAKA_GSTIN = "29AABCU9603R1ZJ"
DELHI_GSTIN = "07AABCU9603R1ZP"

UT_WITHOUT_LEGISLATURE = ["04", "26", "31", "35", "38"]
UT_WITH_LEGISLATURE = ["01", "07", "34"]


def _parse(payload: dict, kind: str) -> dict:
    from apps.engine.parsers.gstr import parse_gstr_file

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "GSTR.json"
        path.write_text(json.dumps(payload), encoding="utf-8")
        return parse_gstr_file(path, kind)


def _gstr2b(supplier_gstin: str, pos, filer_gstin: str = KARNATAKA_GSTIN) -> dict:
    invoice = {
        "inum": "INV-1",
        "dt": "01-07-2026",
        "val": 1180,
        "txval": 1000,
        "iamt": 0,
        "camt": 90,
        "samt": 90,
        "itcavl": "Y",
    }
    if pos is not None:
        invoice["pos"] = pos
    return {
        "gstin": filer_gstin,
        "rtnprd": "072026",
        "data": {
            "docdata": {
                "b2b": [
                    {
                        "ctin": supplier_gstin,
                        "trdnm": "Acme Traders",
                        "inv": [invoice],
                    }
                ]
            }
        },
    }


def _gstr1(filer_gstin: str, pos) -> dict:
    invoice = {
        "inum": "S/1",
        "idt": "01-07-2026",
        "val": 1180,
        "itms": [{"num": 1, "itm_det": {"txval": 1000, "camt": 90, "samt": 90, "csamt": 0}}],
    }
    if pos is not None:
        invoice["pos"] = pos
    return {
        "gstin": filer_gstin,
        "fp": "072026",
        "b2b": [{"ctin": KARNATAKA_GSTIN, "trdnm": "Buyer", "inv": [invoice]}],
    }


class StateMapTests(unittest.TestCase):
    def test_map_covers_every_code_from_01_to_38(self) -> None:
        from apps.engine.validators.states import STATE_NAMES

        expected = [f"{number:02d}" for number in range(1, 39)]
        self.assertEqual(sorted(STATE_NAMES), expected)
        for code, name in STATE_NAMES.items():
            self.assertTrue(name.strip(), code)

    def test_state_names(self) -> None:
        from apps.engine.validators.states import state_name

        self.assertEqual(state_name("04"), "Chandigarh")
        self.assertEqual(state_name("29"), "Karnataka")
        self.assertEqual(state_name(7), "Delhi")
        self.assertIsNone(state_name("99"))
        self.assertIsNone(state_name(None))


class UnionTerritoryRuleTests(unittest.TestCase):
    def test_ut_without_legislature_is_utgst(self) -> None:
        from apps.engine.validators.states import is_union_territory

        for code in UT_WITHOUT_LEGISLATURE:
            self.assertTrue(is_union_territory(code), code)

    def test_ut_with_legislature_is_not_utgst(self) -> None:
        from apps.engine.validators.states import is_union_territory

        for code in UT_WITH_LEGISLATURE:
            self.assertFalse(is_union_territory(code), code)

    def test_states_are_not_utgst(self) -> None:
        from apps.engine.validators.states import is_union_territory

        self.assertFalse(is_union_territory("29"))
        self.assertFalse(is_union_territory("27"))
        self.assertFalse(is_union_territory("99"))
        self.assertFalse(is_union_territory(None))

    def test_intra_ut_helper(self) -> None:
        from apps.engine.validators.states import is_intra_utgst

        self.assertTrue(is_intra_utgst(CHANDIGARH_GSTIN, "04"))
        self.assertTrue(is_intra_utgst(CHANDIGARH_GSTIN, 4))
        self.assertFalse(is_intra_utgst(CHANDIGARH_GSTIN, "29"))
        self.assertFalse(is_intra_utgst(KARNATAKA_GSTIN, "29"))
        self.assertFalse(is_intra_utgst(DELHI_GSTIN, "07"))
        self.assertFalse(is_intra_utgst(None, "04"))


class GstinStateTests(unittest.TestCase):
    def test_state_code_from_real_looking_gstin(self) -> None:
        from apps.engine.validators.states import state_code_from_gstin

        self.assertEqual(state_code_from_gstin(CHANDIGARH_GSTIN), "04")
        self.assertEqual(state_code_from_gstin(KARNATAKA_GSTIN), "29")
        self.assertEqual(state_code_from_gstin("04aabcu9603r1zv"), "04")

    def test_unknown_or_missing_gstin_has_no_state(self) -> None:
        from apps.engine.validators.states import state_code_from_gstin

        self.assertIsNone(state_code_from_gstin("99AABCU9603R1ZZ"))
        self.assertIsNone(state_code_from_gstin(""))
        self.assertIsNone(state_code_from_gstin(None))
        self.assertIsNone(state_code_from_gstin("0"))


class GstrUtgstParserTests(unittest.TestCase):
    def test_intra_chandigarh_2b_row_is_utgst(self) -> None:
        parsed = _parse(_gstr2b(CHANDIGARH_GSTIN, "04"), "gstr_2b")

        row = parsed["rows"][0]
        self.assertIs(row["intra_ut"], True)
        self.assertEqual(row["second_tax_label"], "UTGST")
        self.assertIn("utgst", row["flags"])
        # The portal folds UTGST into the SGST bucket: money stays where it was.
        self.assertEqual(row["sgst"], 90)

    def test_intra_karnataka_2b_row_is_plain_sgst(self) -> None:
        parsed = _parse(_gstr2b(KARNATAKA_GSTIN, "29"), "gstr_2b")

        row = parsed["rows"][0]
        self.assertIs(row["intra_ut"], False)
        self.assertEqual(row["second_tax_label"], "SGST")
        self.assertNotIn("utgst", row["flags"])

    def test_intra_delhi_row_uses_sgst(self) -> None:
        parsed = _parse(_gstr2b(DELHI_GSTIN, "07"), "gstr_2b")

        row = parsed["rows"][0]
        self.assertIs(row["intra_ut"], False)
        self.assertEqual(row["second_tax_label"], "SGST")
        self.assertNotIn("utgst", row["flags"])

    def test_missing_place_of_supply_is_not_classified(self) -> None:
        parsed = _parse(_gstr2b(CHANDIGARH_GSTIN, None), "gstr_2b")

        row = parsed["rows"][0]
        self.assertIs(row["intra_ut"], False)
        self.assertNotIn("utgst", row["flags"])

    def test_gstr1_uses_the_filer_gstin_as_supplier(self) -> None:
        parsed = _parse(_gstr1(CHANDIGARH_GSTIN, "04"), "gstr_1")
        row = next(item for item in parsed["rows"] if item.get("invoice_number") == "S/1")
        self.assertIs(row["intra_ut"], True)
        self.assertEqual(row["second_tax_label"], "UTGST")
        self.assertIn("utgst", row["flags"])

        parsed = _parse(_gstr1(KARNATAKA_GSTIN, "04"), "gstr_1")
        row = next(item for item in parsed["rows"] if item.get("invoice_number") == "S/1")
        self.assertIs(row["intra_ut"], False)
        self.assertNotIn("utgst", row["flags"])

    def test_existing_row_keys_are_untouched(self) -> None:
        parsed = _parse(_gstr2b(CHANDIGARH_GSTIN, "04"), "gstr_2b")

        row = parsed["rows"][0]
        for key in (
            "gstin",
            "trade_name",
            "invoice_number",
            "invoice_date",
            "invoice_value",
            "taxable",
            "igst",
            "cgst",
            "sgst",
            "cess",
            "itc_availability",
            "document_type",
            "hsn",
            "source",
            "flags",
            "match_status",
            "books_ref",
        ):
            self.assertIn(key, row)


class UtgstWorkbookTests(unittest.TestCase):
    def test_b2b_flags_cell_and_cover_note_show_utgst(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.gstr_xlsx import write_gstr_2b
        from apps.engine.pack.style import RED_FILL

        rows = _parse(_gstr2b(CHANDIGARH_GSTIN, "04"), "gstr_2b")["rows"]
        rows += _parse(_gstr2b(KARNATAKA_GSTIN, "29"), "gstr_2b")["rows"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "GSTR_2B_Formatted.xlsx"
            write_gstr_2b(path, rows, {"gstin": CHANDIGARH_GSTIN, "period": "072026"})
            book = load_workbook(path)

        sheet = book["B2B"]
        headers = [cell.value for cell in sheet[1]]
        flags_col = headers.index("Flags") + 1
        flags = [sheet.cell(row, flags_col).value or "" for row in range(2, sheet.max_row + 1)]
        self.assertTrue(any("utgst" in str(value) for value in flags), flags)
        self.assertTrue(any(not value for value in flags), flags)

        red = RED_FILL.fgColor.rgb
        for row in range(2, sheet.max_row + 1):
            fill = sheet.cell(row, 1).fill
            rgb = fill.fgColor.rgb if fill and fill.fgColor is not None else None
            self.assertNotEqual(rgb, red, f"intra-UT row {row} must not look like a defect")

        cover_text = " ".join(
            str(cell.value) for line in book["Cover"].iter_rows() for cell in line if cell.value
        )
        self.assertIn("UTGST", cover_text)
        self.assertIn("SGST column", cover_text)

    def test_table_flags_cell_carries_utgst(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.table_xlsx import write_table

        rows = _parse(_gstr2b(CHANDIGARH_GSTIN, "04"), "gstr_2b")["rows"]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "GSTR_Table.xlsx"
            write_table(
                path,
                "GSTR 2B",
                rows,
                [("GSTIN", "gstin"), ("SGST", "sgst"), ("Flags", "flags")],
            )
            book = load_workbook(path)

        sheet = book["GSTR 2B"]
        self.assertIn("utgst", str(sheet["C2"].value))


if __name__ == "__main__":
    unittest.main()
