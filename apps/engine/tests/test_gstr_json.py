from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

# Valid GSTINs (checksum included) used across the fixtures.
CLIENT_GSTIN = "29ABCDE1234F1ZW"  # Karnataka
MH_PARTY = "27AAPFU0939F1ZV"  # Maharashtra
MH_PARTY_TWO = "27AABCU9603R1ZN"
KA_PARTY = "29AABCU9603R1ZJ"
BROKEN_CHECKSUM = "27AAACN1234A1Z5"

META = {"gstin": CLIENT_GSTIN, "period": "Jul 2026"}
MONEY_KEYS = ("gt", "cur_gt", "val", "txval", "iamt", "camt", "samt", "csamt")


def _row(**values) -> dict:
    row = {
        "source": "books.xml",
        "party_name": "North Retail",
        "invoice_number": "INV-1",
        "invoice_date": "2026-07-08",
    }
    row.update(values)
    return row


def _b2b_row(**values) -> dict:
    row = _row(
        gstin=MH_PARTY,
        taxable_value=1000,
        igst=180,
        cgst=0,
        sgst=0,
        cess=0,
        hsn="9983",
    )
    row.update(values)
    if "invoice_value" not in values and row.get("taxable_value") is not None:
        tax = sum(row.get(key) or 0 for key in ("igst", "cgst", "sgst", "cess"))
        row["invoice_value"] = row["taxable_value"] + tax
    return row


def _walk_money(payload: dict) -> list[float]:
    found: list[float] = []
    for key, value in payload.items():
        if isinstance(value, dict):
            found.extend(_walk_money(value))
        elif isinstance(value, list):
            for item in value:
                if isinstance(item, dict):
                    found.extend(_walk_money(item))
        elif key in MONEY_KEYS and isinstance(value, float):
            found.append(value)
    return found


class GstrJsonRootFieldTests(unittest.TestCase):
    def test_gstr1_root_fields_are_present(self) -> None:
        from apps.engine.pack.gstr_json import GSTR1_VERSION, build_gstr1_json

        payload = build_gstr1_json([_b2b_row()], META)
        for key in ("gstin", "fp", "gt", "cur_gt", "version", "hash"):
            self.assertTrue(payload.get(key) is not None, key)
        self.assertEqual(payload["gstin"], CLIENT_GSTIN)
        self.assertEqual(payload["fp"], "072026")
        self.assertEqual(payload["version"], GSTR1_VERSION)
        self.assertEqual(payload["gt"], 1180.0)
        self.assertEqual(payload["cur_gt"], 1180.0)

    def test_gstr3b_root_fields_are_present(self) -> None:
        from apps.engine.pack.gstr_json import GSTR3B_VERSION, build_gstr3b_json

        payload = build_gstr3b_json([_b2b_row()], META)
        for key in ("gstin", "ret_period", "version", "hash"):
            self.assertTrue(payload.get(key), key)
        self.assertEqual(payload["ret_period"], "072026")
        self.assertEqual(payload["version"], GSTR3B_VERSION)

    def test_period_is_normalised_to_mmyyyy(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json, format_fp

        for raw in ("Jul 2026", "July 2026", "07/2026", "2026-07", "072026", "2026-07-31"):
            self.assertEqual(format_fp(raw), "072026", raw)
        self.assertIsNone(format_fp("last quarter"))
        payload = build_gstr1_json([_b2b_row()], {"gstin": CLIENT_GSTIN, "period": "Nov 2026"})
        self.assertEqual(payload["fp"], "112026")

    def test_missing_gstin_or_period_raises_instead_of_writing_a_file(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json, build_gstr3b_json

        with self.assertRaises(ValueError):
            build_gstr1_json([_b2b_row()], {"period": "Jul 2026"})
        with self.assertRaises(ValueError):
            build_gstr1_json([_b2b_row()], {"gstin": CLIENT_GSTIN})
        with self.assertRaises(ValueError):
            build_gstr3b_json([_b2b_row()], {"gstin": CLIENT_GSTIN})
        with self.assertRaises(ValueError):
            build_gstr1_json([_b2b_row()], {"gstin": BROKEN_CHECKSUM, "period": "Jul 2026"})

    def test_empty_tables_and_doc_issue_are_absent(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json

        payload = build_gstr1_json([_b2b_row(hsn=None)], META)
        self.assertIn("b2b", payload)
        for key in ("b2cs", "b2cl", "cdnr", "cdnur", "hsn", "doc_issue"):
            self.assertNotIn(key, payload, key)


class Gstr1TableTests(unittest.TestCase):
    def test_b2b_groups_by_counterparty_and_keeps_every_line(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json

        rows = [
            _b2b_row(invoice_number="ACME/1", taxable_value=1000, igst=180, hsn="9983"),
            _b2b_row(invoice_number="ACME/1", taxable_value=500, igst=90, hsn="9973"),
            _b2b_row(invoice_number="ACME/2", taxable_value=2000, igst=360, hsn="9983"),
            _b2b_row(invoice_number="OTHER/9", gstin=MH_PARTY_TWO, taxable_value=100, igst=18),
        ]
        payload = build_gstr1_json(rows, META)
        b2b = {party["ctin"]: party for party in payload["b2b"]}
        self.assertEqual(len(b2b), 2)
        invoices = b2b[MH_PARTY]["inv"]
        self.assertEqual([inv["inum"] for inv in invoices], ["ACME/1", "ACME/2"])
        first = invoices[0]
        self.assertEqual(first["idt"], "08-07-2026")
        self.assertEqual(first["val"], 1770.0)
        self.assertEqual(len(first["itms"]), 2)
        self.assertEqual([item["itm_det"]["txval"] for item in first["itms"]], [1000.0, 500.0])
        self.assertEqual([item["itm_det"]["iamt"] for item in first["itms"]], [180.0, 90.0])
        self.assertEqual([item["itm_det"]["rt"] for item in first["itms"]], [18.0, 18.0])
        self.assertEqual(first["itms"][0]["num"], 1)
        self.assertEqual(payload["gt"], 1770.0 + 2360.0 + 118.0)

    def test_b2cs_aggregates_by_rate_and_place_of_supply(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json

        rows = [
            _row(invoice_number="B2C-1", taxable_value=1000, cgst=90, sgst=90, place_of_supply="29-Karnataka"),
            _row(invoice_number="B2C-2", taxable_value=500, cgst=45, sgst=45, place_of_supply="29"),
            _row(invoice_number="B2C-3", taxable_value=1000, cgst=25, sgst=25, place_of_supply="29"),
            _row(invoice_number="B2C-4", taxable_value=1000, igst=180, place_of_supply="27-Maharashtra"),
        ]
        payload = build_gstr1_json(rows, META)
        b2cs = payload["b2cs"]
        self.assertEqual(len(b2cs), 3)
        intra_18 = next(
            item for item in b2cs if item["sply_ty"] == "INTRA" and item["rt"] == 18.0
        )
        self.assertEqual(intra_18["pos"], "29")
        self.assertEqual(intra_18["txval"], 1500.0)
        self.assertEqual(intra_18["camt"], 135.0)
        self.assertEqual(intra_18["samt"], 135.0)
        self.assertEqual(intra_18["typ"], "OE")
        intra_5 = next(item for item in b2cs if item["sply_ty"] == "INTRA" and item["rt"] == 5.0)
        self.assertEqual(intra_5["txval"], 1000.0)
        inter = next(item for item in b2cs if item["sply_ty"] == "INTER")
        self.assertEqual(inter["pos"], "27")
        self.assertEqual(inter["iamt"], 180.0)
        self.assertEqual(inter["camt"], 0.0)
        # Unregistered rows never appear as B2B.
        self.assertNotIn("b2b", payload)

    def test_b2cl_needs_an_inter_state_unregistered_invoice_over_the_threshold(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json

        rows = [
            _row(
                invoice_number="BIG/1",
                taxable_value=300000,
                igst=54000,
                invoice_value=354000,
                place_of_supply="27-Maharashtra",
            )
        ]
        payload = build_gstr1_json(rows, META)
        self.assertIn("b2cl", payload)
        self.assertNotIn("b2cs", payload)
        self.assertEqual(payload["b2cl"][0]["pos"], "27")
        self.assertEqual(payload["b2cl"][0]["inv"][0]["val"], 354000.0)

    def test_credit_note_sign_is_carried_by_ntty(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json, build_gstr3b_json

        rows = [
            _b2b_row(invoice_number="ACME/5", taxable_value=2000, igst=360, invoice_value=2360),
            _b2b_row(
                invoice_number="CN/1",
                document_type="CR",
                taxable_value=-1000,
                igst=-180,
                invoice_value=-1180,
            ),
        ]
        payload = build_gstr1_json(rows, META)
        self.assertEqual(len(payload["cdnr"]), 1)
        note = payload["cdnr"][0]["nt"][0]
        self.assertEqual(note["ntty"], "C")
        self.assertEqual(note["nt_num"], "CN/1")
        self.assertEqual(note["val"], 1180.0)
        self.assertEqual(note["itms"][0]["itm_det"]["txval"], 1000.0)
        self.assertEqual(note["itms"][0]["itm_det"]["iamt"], 180.0)

        summary = build_gstr3b_json(rows, META)["sup_details"]["osup_det"]
        self.assertEqual(summary["txval"], 1000.0)
        self.assertEqual(summary["iamt"], 180.0)


class GstrSplitTests(unittest.TestCase):
    def test_derived_split_is_marked_and_never_outside_the_two_halves_or_igst(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json, build_gstr_report

        rows = [
            _b2b_row(
                gstin=MH_PARTY_TWO,
                taxable_value=2000,
                tax=360,
                igst=None,
                cgst=None,
                sgst=None,
                cess=None,
                tax_rate=18,
            )
        ]
        report = build_gstr_report(rows, META)
        self.assertEqual(report["counts"]["derived_split"], 1)
        derived = report["derived"][0]
        self.assertEqual(derived["split"], "derived")
        self.assertEqual(derived["status"], "derived")
        self.assertEqual(derived["igst"], 360.0)
        self.assertEqual(derived["cgst"], 0.0)

        payload = build_gstr1_json(rows, META)
        item = payload["b2b"][0]["inv"][0]["itms"][0]["itm_det"]
        self.assertEqual(item["iamt"], 360.0)
        self.assertEqual(item["camt"], 0.0)
        self.assertEqual(item["samt"], 0.0)

    def test_intra_state_derived_split_splits_into_equal_halves(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr_report

        rows = [
            _b2b_row(
                gstin=KA_PARTY,
                taxable_value=2000,
                tax=360,
                igst=None,
                cgst=None,
                sgst=None,
                cess=None,
                tax_rate=18,
            )
        ]
        report = build_gstr_report(rows, META)
        entry = report["included"][0]
        self.assertEqual(entry["split"], "derived")
        self.assertEqual(entry["cgst"], 180.0)
        self.assertEqual(entry["sgst"], 180.0)
        self.assertEqual(entry["igst"], 0.0)

    def test_explicit_split_is_not_marked_derived(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr_report

        report = build_gstr_report([_b2b_row()], META)
        self.assertEqual(report["counts"]["derived_split"], 0)
        self.assertEqual(report["included"][0]["split"], "explicit")


class GstrExclusionTests(unittest.TestCase):
    def test_rows_without_a_defensible_split_are_excluded_with_a_reason(self) -> None:
        from apps.engine.pack import gstr_json

        rows = [
            _b2b_row(invoice_number="OK/1"),
            _b2b_row(
                invoice_number="NO-RATE/1",
                tax=180,
                igst=None,
                cgst=None,
                sgst=None,
                cess=None,
            ),
            _b2b_row(invoice_number="BAD-GSTIN/1", gstin=BROKEN_CHECKSUM),
            _row(invoice_number="NO-POS/1", taxable_value=1000, tax=180),
            _row(invoice_number="NO-TAX/1", taxable_value=1000),
        ]
        report = gstr_json.build_gstr_report(rows, META)
        reasons = {item["invoice_number"]: item["reason"] for item in report["excluded"]}
        self.assertEqual(reasons["NO-RATE/1"], gstr_json.REASON_NO_RATE)
        self.assertEqual(reasons["BAD-GSTIN/1"], gstr_json.REASON_BAD_CHECKSUM)
        self.assertEqual(reasons["NO-POS/1"], gstr_json.REASON_NO_POS)
        self.assertEqual(reasons["NO-TAX/1"], gstr_json.REASON_NO_SPLIT)
        self.assertEqual(report["counts"]["excluded"], 4)

        payload = gstr_json.build_gstr1_json(rows, META)
        self.assertEqual(len(payload["b2b"]), 1)
        self.assertEqual(payload["b2b"][0]["ctin"], MH_PARTY)
        self.assertEqual(len(payload["b2b"][0]["inv"]), 1)
        self.assertEqual(payload["gt"], 1180.0)

    def test_nothing_includeable_raises(self) -> None:
        from apps.engine.pack import gstr_json

        rows = [_b2b_row(tax=180, igst=None, cgst=None, sgst=None, cess=None)]
        with self.assertRaises(ValueError):
            gstr_json.build_gstr1_json(rows, META)
        with self.assertRaises(ValueError):
            gstr_json.build_gstr3b_json(rows, META)

    def test_duplicate_rows_are_dropped_once(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json, build_gstr_report

        row = _b2b_row(invoice_number="DUP/1")
        report = build_gstr_report([row, dict(row)], META)
        self.assertEqual(report["counts"]["included"], 1)
        self.assertEqual(report["excluded"][0]["reason"], "identical row already taken from another file")
        payload = build_gstr1_json([row, dict(row)], META)
        self.assertEqual(payload["gt"], 1180.0)


class GstrRoundingTests(unittest.TestCase):
    def test_money_is_rounded_to_paise_everywhere(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr1_json

        rows = [
            _b2b_row(
                taxable_value=1234.567,
                igst=222.2221,
                invoice_value=1456.7891,
                hsn="9983",
            ),
            _row(
                invoice_number="B2C/ROUND",
                taxable_value=333.335,
                cgst=30.0001,
                sgst=30.0001,
                place_of_supply="29",
            ),
        ]
        payload = build_gstr1_json(rows, META)
        item = payload["b2b"][0]["inv"][0]["itms"][0]["itm_det"]
        self.assertEqual(item["txval"], 1234.57)
        self.assertEqual(item["iamt"], 222.22)
        b2cs = payload["b2cs"][0]
        self.assertEqual(b2cs["txval"], 333.34)
        self.assertEqual(b2cs["camt"], 30.0)
        for value in _walk_money(payload):
            self.assertEqual(round(value, 2), value, value)

    def test_derived_halves_always_add_back_to_the_total_tax(self) -> None:
        from apps.engine.pack.gstr_json import build_gstr_report

        rows = [
            _b2b_row(
                gstin=KA_PARTY,
                taxable_value=555.61,
                tax=100.01,
                igst=None,
                cgst=None,
                sgst=None,
                cess=None,
                tax_rate=18,
            )
        ]
        entry = build_gstr_report(rows, META)["included"][0]
        self.assertEqual(round(entry["cgst"] + entry["sgst"], 2), 100.01)
        self.assertEqual(entry["cgst"], 50.01)
        self.assertEqual(entry["sgst"], 50.0)


TALLY_SPLIT_XML = """<?xml version="1.0"?>
<ENVELOPE>
 <HEADER><TALLYREQUEST>Export Data</TALLYREQUEST></HEADER>
 <BODY>
  <TALLYMESSAGE>
   <VOUCHER VCHTYPE="Sales">
    <DATE>20260715</DATE>
    <VOUCHERNUMBER>SAL-IGST</VOUCHERNUMBER>
    <PARTYLEDGERNAME>North Retail</PARTYLEDGERNAME>
    <PARTYGSTIN>27AAPFU0939F1ZV</PARTYGSTIN>
    <PLACEOFSUPPLY>27-Maharashtra</PLACEOFSUPPLY>
    <AMOUNT>1195.00</AMOUNT>
    <TAXABLEVALUE>1000.00</TAXABLEVALUE>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Sales @ 18%</LEDGERNAME>
     <AMOUNT>-1000.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Output IGST @ 18%</LEDGERNAME>
     <AMOUNT>-180.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Output Cess @ 1.5%</LEDGERNAME>
     <AMOUNT>-15.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
   </VOUCHER>
   <VOUCHER VCHTYPE="Sales">
    <DATE>20260716</DATE>
    <VOUCHERNUMBER>SAL-HALVES</VOUCHERNUMBER>
    <PARTYLEDGERNAME>South Spices</PARTYLEDGERNAME>
    <PLACEOFSUPPLY>29-Karnataka</PLACEOFSUPPLY>
    <AMOUNT>1180.00</AMOUNT>
    <TAXABLEVALUE>1000.00</TAXABLEVALUE>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Sales @ 18%</LEDGERNAME>
     <AMOUNT>-1000.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Output CGST @ 9%</LEDGERNAME>
     <AMOUNT>-90.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Output SGST @ 9%</LEDGERNAME>
     <AMOUNT>-90.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
   </VOUCHER>
   <VOUCHER VCHTYPE="Sales">
    <DATE>20260717</DATE>
    <VOUCHERNUMBER>SAL-TAGS</VOUCHERNUMBER>
    <PARTYLEDGERNAME>Tag Traders</PARTYLEDGERNAME>
    <AMOUNT>590.00</AMOUNT>
    <TAXABLEVALUE>500.00</TAXABLEVALUE>
    <CGSTAMOUNT>45.00</CGSTAMOUNT>
    <SGSTAMOUNT>45.00</SGSTAMOUNT>
   </VOUCHER>
  </TALLYMESSAGE>
 </BODY>
</ENVELOPE>
"""


class TallyGstSplitParserTests(unittest.TestCase):
    def _parse(self):
        from apps.engine.parsers.tally import parse_tally_file

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Tally_Sales_Jul2026.xml"
            path.write_text(TALLY_SPLIT_XML, encoding="utf-8")
            return {row["invoice_number"]: row for row in parse_tally_file(path)["rows"]}

    def test_igst_and_cess_come_from_the_ledger_entries(self) -> None:
        rows = self._parse()
        row = rows["SAL-IGST"]
        self.assertEqual(row["igst"], 180.0)
        self.assertEqual(row["cess"], 15.0)
        self.assertIsNone(row["cgst"])
        self.assertIsNone(row["sgst"])
        self.assertEqual(row["tax_rate"], 18.0)
        self.assertEqual(row["place_of_supply"], "27-Maharashtra")
        self.assertEqual(row["place_of_supply_state"], "27")
        # Existing keys are untouched.
        self.assertEqual(row["taxable_value"], 1000.0)
        self.assertEqual(row["register"], "sales")

    def test_cgst_and_sgst_halves_are_read_and_the_rate_adds_up(self) -> None:
        rows = self._parse()
        row = rows["SAL-HALVES"]
        self.assertEqual(row["cgst"], 90.0)
        self.assertEqual(row["sgst"], 90.0)
        self.assertIsNone(row["igst"])
        self.assertEqual(row["tax_rate"], 18.0)
        self.assertEqual(row["place_of_supply_state"], "29")

    def test_voucher_level_tax_tags_are_read(self) -> None:
        rows = self._parse()
        row = rows["SAL-TAGS"]
        self.assertEqual(row["cgst"], 45.0)
        self.assertEqual(row["sgst"], 45.0)
        self.assertIsNone(row["igst"])
        self.assertIsNone(row["place_of_supply_state"])

    def test_voucher_without_tax_ledgers_keeps_none(self) -> None:
        from apps.engine.parsers.tally import parse_tally_file

        xml = """<?xml version="1.0"?>
<ENVELOPE><BODY><TALLYMESSAGE>
 <VOUCHER VCHTYPE="Sales">
  <DATE>20260715</DATE>
  <VOUCHERNUMBER>SAL-PLAIN</VOUCHERNUMBER>
  <PARTYLEDGERNAME>North Retail</PARTYLEDGERNAME>
  <AMOUNT>5900.00</AMOUNT>
 </VOUCHER>
</TALLYMESSAGE></BODY></ENVELOPE>
"""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "plain.xml"
            path.write_text(xml, encoding="utf-8")
            row = parse_tally_file(path)["rows"][0]
        for key in ("igst", "cgst", "sgst", "cess", "tax_rate", "place_of_supply_state"):
            self.assertIsNone(row[key], key)


class ZohoGstSplitParserTests(unittest.TestCase):
    def _parse(self, headers: list[str], values: list[list]):
        from openpyxl import Workbook

        from apps.engine.parsers.zoho import parse_zoho_file

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "zoho_invoices.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.append(headers)
            for row in values:
                sheet.append(row)
            book.save(path)
            return parse_zoho_file(path)

    def test_cgst_sgst_columns_and_place_of_supply(self) -> None:
        parsed = self._parse(
            [
                "Invoice Number",
                "Invoice Date",
                "GST Identification Number (GSTIN)",
                "Place of Supply",
                "Subtotal",
                "CGST",
                "SGST",
                "Total",
                "HSN/SAC",
            ],
            [["INV-9", "2026-07-10", KA_PARTY, "29-Karnataka", 1000, 90, 90, 1180, "9983"]],
        )
        row = parsed["rows"][0]
        self.assertEqual(row["cgst"], 90.0)
        self.assertEqual(row["sgst"], 90.0)
        self.assertIsNone(row["igst"])
        self.assertIsNone(row["cess"])
        # No single "tax" column: the heads add up instead.
        self.assertEqual(row["tax"], 180.0)
        self.assertEqual(row["place_of_supply"], "29-Karnataka")
        self.assertEqual(row["place_of_supply_state"], "29")
        self.assertEqual(row["taxable_value"], 1000.0)
        self.assertEqual(row["register"], "sales")

    def test_igst_and_cess_columns(self) -> None:
        parsed = self._parse(
            ["Invoice Number", "Invoice Date", "GSTIN", "Subtotal", "IGST", "Cess Amount", "Total"],
            [["INV-10", "2026-07-11", MH_PARTY, 2000, 360, 24, 2384]],
        )
        row = parsed["rows"][0]
        self.assertEqual(row["igst"], 360.0)
        self.assertEqual(row["cess"], 24.0)
        self.assertEqual(row["tax"], 384.0)
        self.assertIsNone(row["cgst"])

    def test_tax_column_still_wins_over_the_heads(self) -> None:
        parsed = self._parse(
            ["Invoice Number", "Subtotal", "CGST", "SGST", "Tax Amount", "IGST"],
            [["INV-11", 1000, 90, 90, 180, 0]],
        )
        row = parsed["rows"][0]
        self.assertEqual(row["tax"], 180.0)
        self.assertEqual(row["igst"], 0.0)


TALLY_PACK_XML = """<?xml version="1.0"?>
<ENVELOPE>
 <HEADER><TALLYREQUEST>Export Data</TALLYREQUEST></HEADER>
 <BODY>
  <TALLYMESSAGE>
   <VOUCHER VCHTYPE="Sales">
    <DATE>20260715</DATE>
    <VOUCHERNUMBER>SAL-CLEAN</VOUCHERNUMBER>
    <PARTYLEDGERNAME>North Retail</PARTYLEDGERNAME>
    <PARTYGSTIN>27AAPFU0939F1ZV</PARTYGSTIN>
    <AMOUNT>1180.00</AMOUNT>
    <TAXABLEVALUE>1000.00</TAXABLEVALUE>
    <HSNCODE>9983</HSNCODE>
    <ALLLEDGERENTRIES.LIST>
     <LEDGERNAME>Output IGST @ 18%</LEDGERNAME>
     <AMOUNT>-180.00</AMOUNT>
    </ALLLEDGERENTRIES.LIST>
   </VOUCHER>
   <VOUCHER VCHTYPE="Sales">
    <DATE>20260716</DATE>
    <VOUCHERNUMBER>SAL-NO-RATE</VOUCHERNUMBER>
    <PARTYLEDGERNAME>Mystery Buyer</PARTYLEDGERNAME>
    <PARTYGSTIN>27AABCU9603R1ZN</PARTYGSTIN>
    <AMOUNT>1180.00</AMOUNT>
    <TAXABLEVALUE>1000.00</TAXABLEVALUE>
    <TAXAMOUNT>180.00</TAXAMOUNT>
   </VOUCHER>
  </TALLYMESSAGE>
 </BODY>
</ENVELOPE>
"""


class GstrJsonPackTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._tmp.cleanup()

    def _pack(self, gstin: str | None) -> tuple[dict, Path]:
        from apps.engine.clients import create_client
        from apps.engine.dump import ingest_paths, start_job
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period
        from apps.engine.pipeline import get_period_pack

        source = Path(self._tmp.name) / "Tally_Sales_Jul2026.xml"
        source.write_text(TALLY_PACK_XML, encoding="utf-8")

        save_firm("Test firm")
        client = create_client("Acme", gstin=gstin)
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(source)])
        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        return pack, Path(pack["path"])

    def test_pack_gets_upload_json_and_the_cannot_include_sheet(self) -> None:
        pack, folder = self._pack(CLIENT_GSTIN)
        by_key = {item["key"]: item for item in pack["outputs"]}
        self.assertIn("sales", by_key)
        self.assertIn("gstr_1_json", by_key)
        self.assertIn("gstr_3b_json", by_key)
        self.assertIn("gstr_cannot_include", by_key)

        gstr1_path = Path(by_key["gstr_1_json"]["path"])
        self.assertEqual(gstr1_path.name, "GSTR_1_Upload.json")
        payload = json.loads(gstr1_path.read_text(encoding="utf-8"))
        self.assertEqual(payload["gstin"], CLIENT_GSTIN)
        self.assertEqual(payload["fp"], "072026")
        self.assertEqual(payload["version"], "GST3.2.0")
        self.assertEqual(payload["b2b"][0]["ctin"], MH_PARTY)
        self.assertEqual(payload["b2b"][0]["inv"][0]["inum"], "SAL-CLEAN")

        gstr3b_path = Path(by_key["gstr_3b_json"]["path"])
        summary_3b = json.loads(gstr3b_path.read_text(encoding="utf-8"))
        self.assertEqual(summary_3b["ret_period"], "072026")
        self.assertEqual(summary_3b["sup_details"]["osup_det"]["iamt"], 180.0)

        cannot = Path(by_key["gstr_cannot_include"]["path"])
        book = load_workbook(cannot)
        self.assertIn("Cannot include", book.sheetnames)
        texts = [
            str(cell)
            for row in book["Cannot include"].iter_rows(min_row=2, values_only=True)
            for cell in row
            if cell is not None
        ]
        joined = " ".join(texts)
        self.assertIn("SAL-NO-RATE", joined)
        self.assertIn("tax rate", joined)

        summary = json.loads((folder / "pack_summary.json").read_text(encoding="utf-8"))
        block = summary["gstr_json"]
        self.assertEqual(block["files"]["gstr_1_json"]["status"], "written")
        self.assertEqual(block["counts"]["included"], 1)
        self.assertEqual(block["counts"]["excluded"], 1)
        self.assertEqual(len(block["excluded"]), 1)

    def test_without_a_client_gstin_nothing_is_written_and_the_reason_is_recorded(self) -> None:
        pack, folder = self._pack(None)
        by_key = {item["key"]: item for item in pack["outputs"]}
        self.assertIn("sales", by_key)
        self.assertNotIn("gstr_1_json", by_key)
        self.assertNotIn("gstr_3b_json", by_key)
        self.assertFalse((folder / "GSTR_1_Upload.json").exists())
        self.assertFalse((folder / "GSTR_3B_Upload.json").exists())

        summary = json.loads((folder / "pack_summary.json").read_text(encoding="utf-8"))
        block = summary["gstr_json"]
        self.assertEqual(block["files"]["gstr_1_json"]["status"], "skipped")
        self.assertIn("GSTIN", block["files"]["gstr_1_json"]["reason"])
        self.assertIn("GSTIN", block["files"]["gstr_3b_json"]["reason"])


if __name__ == "__main__":
    unittest.main()
