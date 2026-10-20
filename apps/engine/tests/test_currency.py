from __future__ import annotations

import os
import sys
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from make_test_dump import PNG, pdf_with_text


def pdf_pages(pages: list[list[str]]) -> bytes:
    """Build a small multi-page PDF with one text block per page."""
    objects: list[bytes] = []
    page_count = len(pages)
    font_number = 3 + 2 * page_count
    kids = " ".join(f"{3 + 2 * index} 0 R" for index in range(page_count))
    objects.append(b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n")
    objects.append(
        f"2 0 obj<</Type/Pages/Kids[{kids}]/Count {page_count}>>endobj\n".encode()
    )
    for index, lines in enumerate(pages):
        content = "BT /F1 10 Tf 40 800 Td 14 TL\n"
        for line_index, line in enumerate(lines):
            safe = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            if line_index == 0:
                content += f"({safe}) Tj T*\n"
            else:
                content += f"({safe}) '\n"
        content += "ET\n"
        stream = content.encode("latin-1", errors="replace")
        page_number = 3 + 2 * index
        content_number = 4 + 2 * index
        objects.append(
            (
                f"{page_number} 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 595 842]"
                f"/Contents {content_number} 0 R/Resources<</Font<</F1 {font_number} 0 R>>>>>>"
                "endobj\n"
            ).encode()
        )
        objects.append(
            f"{content_number} 0 obj<</Length {len(stream)}>>stream\n".encode()
            + stream
            + b"endstream\nendobj\n"
        )
    objects.append(
        f"{font_number} 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj\n".encode()
    )
    body = b"%PDF-1.4\n"
    offsets = [0]
    for obj in objects:
        offsets.append(len(body))
        body += obj
    xref_position = len(body)
    xref = f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    for offset in offsets[1:]:
        xref += f"{offset:010d} 00000 n \n"
    trailer = (
        f"trailer<</Size {len(objects) + 1}/Root 1 0 R>>\n"
        f"startxref\n{xref_position}\n%%EOF\n"
    )
    return body + xref.encode() + trailer.encode()


def large_statement_pages(page_count: int = 25, per_page: int = 20) -> list[list[str]]:
    balance = 250_000.0
    pages: list[list[str]] = []
    for page in range(page_count):
        lines: list[str] = []
        if page == 0:
            lines.extend(
                [
                    "HDFC Bank Account Statement",
                    "Account Number 50100123456789",
                    "IFSC HDFC0001234",
                    "Opening Balance 250000.00",
                ]
            )
        for offset in range(per_page):
            index = page * per_page + offset
            if index % 2 == 0:
                balance -= 125.0
                lines.append(
                    f"01/07/2026 TXN {index:05d} Withdrawal 125.00 {balance:.2f}"
                )
            else:
                balance += 75.0
                lines.append(f"01/07/2026 TXN {index:05d} Deposit 75.00 {balance:.2f}")
        if page == page_count - 1:
            lines.append(f"Closing Balance {balance:.2f}")
        pages.append(lines)
    return pages


class CurrencyRegistryTests(unittest.TestCase):
    def test_registry_is_broad_and_offline(self) -> None:
        from apps.engine.currency import active_currency_codes, currency_options

        codes = set(active_currency_codes())
        self.assertGreaterEqual(len(codes), 150)
        for expected in {"INR", "USD", "EUR", "GBP", "JPY", "AED", "SGD", "ZAR"}:
            self.assertIn(expected, codes)
        options = currency_options()
        self.assertGreaterEqual(len(options), 150)
        self.assertTrue(all(option["code"] and option["label"] for option in options))

    def test_historical_codes_are_known_but_not_active(self) -> None:
        from apps.engine.currency import is_known, is_supported

        self.assertTrue(is_known("DEM"))
        self.assertFalse(is_supported("DEM"))
        self.assertFalse(is_known("ZZZ"))

    def test_currency_labels_are_explicit(self) -> None:
        from apps.engine.currency import currency_label, currency_options

        self.assertEqual(
            currency_label({"currency_code": "USD", "currency_status": "confirmed"}),
            "USD",
        )
        self.assertEqual(
            currency_label({"currency_code": "USD", "currency_status": "inferred"}),
            "USD (inferred)",
        )
        self.assertIn("Unknown", currency_label({"currency_status": "ambiguous"}))
        self.assertTrue(any(option["code"] == "USD" for option in currency_options()))

    def test_excel_formats_follow_currency_precision(self) -> None:
        from apps.engine.currency import excel_number_format, minor_units

        self.assertEqual(minor_units("JPY"), 0)
        self.assertEqual(minor_units("KWD"), 3)
        self.assertNotIn(".00", excel_number_format("JPY"))
        self.assertIn("USD", excel_number_format("USD"))
        self.assertIn("(inferred)", excel_number_format("USD", inferred=True))


class CurrencyDetectionTests(unittest.TestCase):
    def test_labelled_currency_is_confirmed(self) -> None:
        from apps.engine.currency import detect_currency

        decision = detect_currency("Currency: USD\nOpening Balance 1,000.00")
        self.assertEqual(decision["currency_code"], "USD")
        self.assertEqual(decision["currency_status"], "confirmed")

    def test_iso_code_next_to_amount_is_confirmed(self) -> None:
        from apps.engine.currency import detect_currency

        decision = detect_currency("Opening balance 1.234,56 EUR")
        self.assertEqual(decision["currency_code"], "EUR")
        self.assertEqual(decision["currency_status"], "confirmed")

    def test_bare_dollar_stays_ambiguous(self) -> None:
        from apps.engine.currency import detect_currency

        decision = detect_currency("Opening Balance $ 1,000.00")
        self.assertIsNone(decision["currency_code"])
        self.assertEqual(decision["currency_status"], "ambiguous")
        self.assertTrue(decision["currency_ambiguous"])

    def test_country_metadata_resolves_bare_dollar(self) -> None:
        from apps.engine.currency import detect_currency

        decision = detect_currency("United States account\nBalance $ 1,000.00")
        self.assertEqual(decision["currency_code"], "USD")
        self.assertEqual(decision["currency_status"], "inferred")
        self.assertEqual(decision["currency_source"], "ambiguous_symbol+country_metadata")

    def test_conflicting_explicit_codes_are_ambiguous(self) -> None:
        from apps.engine.currency import detect_currency

        decision = detect_currency("Currency: USD\nCurrency: EUR")
        self.assertIsNone(decision["currency_code"])
        self.assertEqual(decision["currency_status"], "ambiguous")

    def test_ifsc_is_only_weak_inr_evidence(self) -> None:
        from apps.engine.currency import detect_currency

        decision = detect_currency("IFSC HDFC0001234")
        self.assertEqual(decision["currency_code"], "INR")
        self.assertEqual(decision["currency_status"], "inferred")
        self.assertEqual(decision["currency_source"], "bank_metadata")

        dollar = detect_currency("IFSC HDFC0001234\nBalance $ 1,000.00")
        self.assertIsNone(dollar["currency_code"])
        self.assertEqual(dollar["currency_status"], "ambiguous")

    def test_apply_override_validates_iso_codes(self) -> None:
        from apps.engine.currency import apply_override

        decision = {
            "currency_code": None,
            "currency_status": "ambiguous",
            "currency_ambiguous": True,
        }
        updated = apply_override(decision, "USD")
        self.assertEqual(updated["currency_code"], "USD")
        self.assertEqual(updated["currency_status"], "confirmed")
        self.assertEqual(updated["currency_source"], "manual_override")
        with self.assertRaisesRegex(ValueError, "ISO currency"):
            apply_override(decision, "NOPE")


class InternationalMoneyTests(unittest.TestCase):
    def test_amount_conventions(self) -> None:
        from apps.engine.parsers.bank.money import parse_amount

        self.assertEqual(parse_amount("1,234.56"), Decimal("1234.56"))
        self.assertEqual(parse_amount("1.234,56"), Decimal("1234.56"))
        self.assertEqual(parse_amount("1 234,56"), Decimal("1234.56"))
        self.assertEqual(parse_amount("1,234"), Decimal("1234"))
        self.assertIsNone(parse_amount("1.234"))
        self.assertEqual(parse_amount("1.234", decimal_separator=","), Decimal("1234"))
        self.assertEqual(parse_amount("1,234", decimal_separator="."), Decimal("1234"))
        self.assertEqual(parse_amount("(1,234.56)"), Decimal("-1234.56"))
        self.assertEqual(parse_amount("1,234.56 Dr"), Decimal("1234.56"))

    def test_currency_tokens_are_ignored_for_value(self) -> None:
        from apps.engine.parsers.bank.money import parse_amount

        self.assertEqual(parse_amount("₹1,23,456.78"), Decimal("123456.78"))
        self.assertEqual(parse_amount("€ 1.234,56"), Decimal("1234.56"))
        self.assertEqual(parse_amount("USD 99.50"), Decimal("99.50"))

    def test_decimal_separator_detection(self) -> None:
        from apps.engine.parsers.bank.money import detect_decimal_separator

        self.assertEqual(detect_decimal_separator("Opening 1.234,56"), ",")
        self.assertEqual(detect_decimal_separator("Opening 1,234.56"), ".")

    def test_date_fragments_are_not_amounts(self) -> None:
        from apps.engine.parsers.bank.money import amounts_in

        self.assertEqual(amounts_in("01.07.2026 UPI 2,500.00"), [Decimal("2500.00")])
        self.assertEqual(amounts_in("Txn on 01.07.2026"), [])


class ImageStatementTests(unittest.TestCase):
    def _extracted(self):
        from apps.engine.pdf_extract import ExtractedPdf, LineBox, PageText

        text_lines = [
            "Bank Account Statement",
            "Opening Balance 1000.00",
            "01/07/2026 Deposit 100.00 1100.00",
            "Closing Balance 1100.00",
        ]
        lines = [
            LineBox(page=1, text=value, x=20, y=700 - index * 20, width=300, height=12)
            for index, value in enumerate(text_lines)
        ]
        return ExtractedPdf(
            pdf_type="image_based",
            page_count=1,
            pages=[PageText(page=1, text="\n".join(text_lines), needs_ocr=False)],
            lines=lines,
            words=[],
            engine="rapidocr",
        )

    def test_jpg_jpeg_and_png_use_the_bank_parser(self) -> None:
        from apps.engine.parsers.bank.parser import parse_bank_statement
        from apps.engine.validators.balance import check_balance

        with tempfile.TemporaryDirectory() as tmp:
            for suffix in (".png", ".jpg", ".jpeg"):
                path = Path(tmp) / f"statement{suffix}"
                path.write_bytes(PNG)
                with patch("apps.engine.ocr.extract_image", return_value=self._extracted()):
                    parsed = parse_bank_statement(path)
                self.assertEqual(parsed["input_type"], "image", suffix)
                self.assertEqual(parsed["engine"], "rapidocr", suffix)
                self.assertEqual(parsed["parse_strategy"], "lines", suffix)
                self.assertEqual(len(parsed["rows"]), 1, suffix)
                check = check_balance(
                    parsed["rows"], parsed["opening_balance"], parsed["stated_closing"]
                )
                self.assertEqual(check["status"], "match", suffix)

    def test_exif_orientation_is_applied(self) -> None:
        from apps.engine.ocr import _prepare_image

        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "rotated.jpg"
            image = Image.new("RGB", (40, 20), "white")
            exif = Image.Exif()
            exif[274] = 6  # rotate 90 degrees clockwise
            image.save(source, "JPEG", exif=exif)
            prepared = _prepare_image(source, Path(tmp))
            self.assertIsNotNone(prepared)
            assert prepared is not None
            _, width, height, _scale = prepared
            self.assertEqual((width, height), (20.0, 40.0))

    def test_classifier_uses_ocr_text_for_images(self) -> None:
        from apps.engine.classifier import classify_path

        from PIL import Image

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scan.png"
            Image.new("RGB", (400, 200), "white").save(path)
            with patch("apps.engine.classifier._ocr_available", return_value=True), patch(
                "apps.engine.classifier._ocr_image_text",
                return_value=(
                    "Bank Account Statement\nOpening Balance 1000.00\n"
                    "Deposit 100.00\nClosing Balance 1100.00"
                ),
            ):
                result = classify_path(path)
            self.assertEqual(result.kind, "bank", result.reason)

    def test_scanned_pdf_uses_ocr_text_instead_of_filename(self) -> None:
        from apps.engine.classifier import classify_path
        from apps.engine.pdf_extract import ExtractedPdf, PageText

        extracted = ExtractedPdf(
            pdf_type="scanned",
            page_count=4,
            pages=[
                PageText(
                    page=1,
                    text=(
                        "Bank Account Statement\nOpening Balance 100.00\n"
                        "Withdrawal 10.00\nDeposit 5.00\nClosing Balance 95.00"
                    ),
                    needs_ocr=False,
                )
            ],
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "scan_004.pdf"
            path.write_bytes(b"%PDF-1.4\n")
            with patch("apps.engine.pdf_extract.extract_pdf", return_value=extracted):
                result = classify_path(path)
        self.assertEqual(result.kind, "bank", result.reason)


class WorkbookCurrencyTests(unittest.TestCase):
    def _rows(self, currency_code="USD", status="confirmed"):
        return [
            {
                "date": "2026-07-01",
                "description": "Opening deposit",
                "cheque_ref": None,
                "debit": None,
                "credit": 1000.00,
                "balance": 1000.00,
                "source": "statement.pdf#p1",
                "flags": [],
                "currency_code": currency_code,
                "currency_status": status,
            },
            {
                "date": "2026-07-02",
                "description": "Withdrawal",
                "cheque_ref": None,
                "debit": 100.00,
                "credit": None,
                "balance": 900.00,
                "source": "statement.pdf#p1",
                "flags": ["inferred_from_balance"],
                "currency_code": currency_code,
                "currency_status": status,
            },
        ]

    def test_currency_columns_formats_and_removed_instruction(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.bank_xlsx import COVER_HEADERS, HEADERS, write_bank_workbook
        from apps.engine.validators.balance import check_balance

        rows = self._rows()
        check = check_balance(rows, 0.0, 900.0)
        currency = {
            "currency_code": "USD",
            "currency_status": "confirmed",
            "currency_source": "labelled_field",
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bank.xlsx"
            write_bank_workbook(
                path,
                [
                    {
                        "rows": rows,
                        "check": check,
                        "meta": {"filename": "statement.pdf", "profile_label": "HDFC"},
                        "currency": currency,
                    }
                ],
            )
            book = load_workbook(path)
            cover = book["Balance Check"]
            sheet = book["Transactions"]
            cover_headers = [cell.value for cell in cover[4]]
            tx_headers = [cell.value for cell in sheet[1]]
            self.assertEqual(cover_headers, COVER_HEADERS)
            self.assertEqual(tx_headers, HEADERS)
            self.assertEqual(cover_headers[-1], "Currency")
            self.assertEqual(tx_headers[-1], "Currency")
            self.assertEqual(cover.cell(5, len(COVER_HEADERS)).value, "USD")
            self.assertEqual(sheet.cell(2, len(HEADERS)).value, "USD")
            self.assertIn("USD", sheet.cell(2, 4).number_format)
            self.assertIn("USD", sheet.cell(2, 6).number_format)
            self.assertEqual(sheet.cell(3, len(HEADERS)).value, "USD")
            removed = (
                "Each source file is checked on its own. "
                "Do not mix two statements when you test."
            )
            values = [
                str(cell.value)
                for row in cover.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            values += [
                str(cell.value)
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            ]
            self.assertNotIn(removed, values)

    def test_inferred_currency_is_visible_in_formats(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.bank_xlsx import write_bank_workbook
        from apps.engine.validators.balance import check_balance

        rows = self._rows(status="inferred")
        check = check_balance(rows, 0.0, 900.0)
        currency = {
            "currency_code": "USD",
            "currency_status": "inferred",
            "currency_source": "country_metadata",
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "bank.xlsx"
            write_bank_workbook(
                path,
                [
                    {
                        "rows": rows,
                        "check": check,
                        "meta": {"filename": "statement.pdf"},
                        "currency": currency,
                    }
                ],
            )
            sheet = load_workbook(path)["Transactions"]
            self.assertEqual(sheet.cell(2, len(sheet[1])).value, "USD (inferred)")
            self.assertIn("(inferred)", sheet.cell(2, 4).number_format)
            self.assertIn("(inferred)", sheet.cell(2, 6).number_format)

    def test_workbook_capacity_guard_never_truncates(self) -> None:
        from apps.engine.pack.bank_xlsx import MAX_BANK_SHEET_ROWS, write_bank_workbook

        class HugeRows:
            def __len__(self) -> int:
                return MAX_BANK_SHEET_ROWS + 1

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "too-large.xlsx"
            with self.assertRaisesRegex(ValueError, "No rows were written"):
                write_bank_workbook(
                    path,
                    [
                        {
                            "rows": HugeRows(),
                            "check": {"row_count": MAX_BANK_SHEET_ROWS + 1},
                            "meta": {"filename": "huge.pdf"},
                            "currency": {},
                        }
                    ],
                )
            self.assertFalse(path.exists())


class LargeStatementIntegrityTests(unittest.TestCase):
    def test_multi_page_statement_keeps_every_transaction(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.bank_xlsx import write_bank_workbook
        from apps.engine.parsers.bank.parser import parse_bank_statement
        from apps.engine.validators.balance import check_balance

        page_count, per_page = 25, 20
        total = page_count * per_page
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "Large_Statement.pdf"
            path.write_bytes(pdf_pages(large_statement_pages(page_count, per_page)))
            parsed = parse_bank_statement(path)
            self.assertEqual(len(parsed["rows"]), total)
            self.assertEqual(parsed["candidate_count"], total)
            self.assertEqual(parsed["dropped_count"], 0)
            self.assertEqual(parsed["duplicate_count"], 0)
            self.assertEqual(parsed["inferred_count"], 0)
            check = check_balance(
                parsed["rows"], parsed["opening_balance"], parsed["stated_closing"]
            )
            self.assertEqual(check["status"], "match", check)
            indexes = set()
            for row in parsed["rows"]:
                description = row["description"]
                marker = description.split("TXN ", 1)[1][:5]
                indexes.add(int(marker))
            self.assertEqual(indexes, set(range(total)))

            dest = Path(tmp) / "large.xlsx"
            write_bank_workbook(
                dest,
                [
                    {
                        "rows": parsed["rows"],
                        "check": check,
                        "meta": {"filename": path.name},
                        "currency": parsed["currency"],
                    }
                ],
            )
            sheet = load_workbook(dest, read_only=True)["Transactions"]
            try:
                self.assertEqual(sheet.max_row, total + 1)
            finally:
                sheet.parent.close()

    def test_pdf_page_guard_stops_before_extraction(self) -> None:
        from pypdf import PdfWriter

        from apps.engine.parsers.bank.parser import MAX_BANK_PDF_PAGES, parse_bank_statement

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "too-many-pages.pdf"
            writer = PdfWriter()
            for _ in range(MAX_BANK_PDF_PAGES + 1):
                writer.add_blank_page(width=595, height=842)
            with path.open("wb") as handle:
                writer.write(handle)
            with self.assertRaisesRegex(ValueError, "safe bank-statement limit"):
                parse_bank_statement(path)

    def test_transaction_limit_raises_instead_of_truncating(self) -> None:
        from apps.engine.parsers.bank.parser import parse_bank_statement
        from apps.engine.tests.test_stage3_bank import write_hdfc

        with tempfile.TemporaryDirectory() as tmp:
            path = write_hdfc(Path(tmp))
            with patch(
                "apps.engine.parsers.bank.parser.MAX_BANK_TRANSACTIONS", 2
            ):
                with self.assertRaisesRegex(ValueError, "Split the statement"):
                    parse_bank_statement(path)

    def test_duplicate_source_rows_are_flagged_not_removed(self) -> None:
        from apps.engine.parsers.bank.parser import _flag_duplicate_source_rows

        rows = [
            {
                "date": "2026-07-01",
                "description": "UPI merchant",
                "debit": 100.0,
                "credit": None,
                "balance": 900.0,
                "source": "statement.pdf#p1@1,2,3,4",
                "source_page": 1,
                "source_bbox": "1,2,3,4",
                "flags": [],
            },
            {
                "date": "2026-07-01",
                "description": "UPI merchant",
                "debit": 100.0,
                "credit": None,
                "balance": 900.0,
                "source": "statement.pdf#p1@1,2,3,4",
                "source_page": 1,
                "source_bbox": "1,2,3,4",
                "flags": [],
            },
        ]
        self.assertEqual(_flag_duplicate_source_rows(rows), 1)
        self.assertEqual(len(rows), 2)
        self.assertIn("duplicate_source_row", rows[1]["flags"])
        self.assertNotIn("duplicate_source_row", rows[0]["flags"])


class CurrencyPersistenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self._env = patch.dict(os.environ, {"LOCALAPPDATA": self._tmp.name})
        self._env.start()
        from apps.engine.db import reset_engine

        reset_engine()

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._env.stop()
        self._tmp.cleanup()

    def _period(self):
        from apps.engine.clients import create_client
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period

        save_firm("Test firm")
        client = create_client("Acme")
        return create_period(client["id"], "Jul 2026")

    def _ingest(self, period, filename: str, lines: list[str]):
        from apps.engine.dump import ingest_paths, list_period_files, start_job

        inbox = Path(self._tmp.name) / f"inbox-{filename}"
        inbox.mkdir()
        path = inbox / filename
        path.write_bytes(pdf_with_text(lines))
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(path)])
        return next(
            row for row in list_period_files(period["id"]) if row["original_name"] == filename
        )

    def test_confirmed_currency_is_persisted_and_override_is_guarded(self) -> None:
        from apps.engine.dump import set_file_currency
        from apps.engine.pipeline import get_period_pack

        period = self._period()
        file_row = self._ingest(
            period,
            "USD_Statement_Jul2026.pdf",
            [
                "USD Bank Account Statement",
                "Currency: USD",
                "Account Number 50100123456789",
                "IFSC HDFC0001234",
                "Opening Balance 1,000.00",
                "01/07/2026 Deposit 100.00 1,100.00",
                "Closing Balance 1,100.00",
            ],
        )
        self.assertEqual(file_row["currency_code"], "USD")
        self.assertEqual(file_row["currency_status"], "confirmed")
        self.assertEqual(file_row["currency_label"], "USD")

        with self.assertRaisesRegex(ValueError, "clearly shows USD"):
            set_file_currency(file_row["id"], "EUR")

        same = set_file_currency(file_row["id"], "USD")
        self.assertEqual(same["currency_override"], "USD")
        self.assertEqual(same["currency_label"], "USD")

        cleared = set_file_currency(file_row["id"], "")
        self.assertIsNone(cleared["currency_override"])
        self.assertEqual(cleared["currency_code"], "USD")
        self.assertEqual(cleared["currency_status"], "confirmed")

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        assert pack is not None
        self.assertEqual(pack["files"][0]["currency_code"], "USD")
        self.assertEqual(pack["files"][0]["currency_label"], "USD")

    def test_ambiguous_currency_can_be_overridden_and_reparsed(self) -> None:
        from apps.engine.dump import list_period_files, reparse_period, set_file_currency
        from apps.engine.pipeline import get_period_pack, get_period_preview

        period = self._period()
        file_row = self._ingest(
            period,
            "Ambiguous_Jul2026.pdf",
            [
                "Bank Account Statement",
                "Account Number 50100123456789",
                "Opening Balance $ 1,000.00",
                "01/07/2026 Deposit 100.00 1,100.00",
                "Closing Balance 1,100.00",
            ],
        )
        self.assertIsNone(file_row["currency_code"])
        self.assertEqual(file_row["currency_status"], "ambiguous")
        self.assertIn("Unknown", file_row["currency_label"])

        updated = set_file_currency(file_row["id"], "USD")
        self.assertEqual(updated["currency_override"], "USD")
        self.assertEqual(updated["currency_code"], "USD")
        self.assertEqual(updated["currency_status"], "confirmed")
        self.assertEqual(updated["currency_source"], "manual_override")

        parsed = reparse_period(period["id"])
        self.assertEqual(parsed["status"], "done")
        refreshed = next(
            row
            for row in list_period_files(period["id"])
            if row["original_name"] == "Ambiguous_Jul2026.pdf"
        )
        self.assertEqual(refreshed["currency_override"], "USD")
        self.assertEqual(refreshed["currency_code"], "USD")
        self.assertEqual(refreshed["currency_label"], "USD")

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        assert pack is not None
        self.assertEqual(pack["files"][0]["currency_code"], "USD")
        self.assertEqual(pack["files"][0]["currency_source"], "manual_override")

        preview = get_period_preview(period["id"])
        bank = next(file for file in preview["files"] if file["kind"] == "bank")
        self.assertEqual(bank["currency"]["currency_code"], "USD")
        self.assertEqual(bank["currency"]["currency_label"], "USD")
        self.assertEqual(bank["preview"][0]["currency_code"], "USD")

    def test_non_inr_rows_are_filtered_from_reconciliation(self) -> None:
        from apps.engine.pipeline import _recon_bank_rows

        bank_files = [
            {
                "rows": [
                    {"description": "usd", "currency_code": "USD"},
                    {"description": "inr", "currency_code": "INR"},
                    {"description": "unknown", "currency_code": None},
                ]
            }
        ]
        rows = _recon_bank_rows(bank_files)
        self.assertEqual([row["description"] for row in rows], ["inr", "unknown"])


if __name__ == "__main__":
    unittest.main()
