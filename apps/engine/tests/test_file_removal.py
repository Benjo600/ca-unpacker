from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from apps.engine.tests.test_stage3_bank import write_hdfc


class RemovePeriodFileTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()

        from apps.engine.clients import create_client
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period

        save_firm("Test firm")
        self.client = create_client("Acme")
        self.period = create_period(self.client["id"], "Jul 2026")
        self.inbox = Path(self._tmp.name) / "inbox"
        self.inbox.mkdir()

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._tmp.cleanup()

    def _two_bank_files(self) -> tuple[Path, Path]:
        first = write_hdfc(self.inbox)
        first = first.rename(self.inbox / "HDFC_Statement_A.pdf")
        second = write_hdfc(self.inbox)
        second = second.rename(self.inbox / "HDFC_Statement_B.pdf")
        return first, second

    def _ingest(self, *paths: Path) -> dict:
        from apps.engine.dump import ingest_paths, start_job

        job = start_job(self.period["id"])
        return ingest_paths(job["id"], [str(path) for path in paths])

    def test_removing_one_file_keeps_the_other_file_and_its_rows(self) -> None:
        from apps.engine.db import ExtractedRow, StoredFile, get_session
        from apps.engine.dump import list_period_files
        from apps.engine.pipeline import get_period_preview
        from apps.engine.removal import remove_period_file

        first, second = self._two_bank_files()
        job = self._ingest(first, second)
        self.assertEqual(len(job["files"]), 2)
        files_before = {item["original_name"]: item["id"] for item in list_period_files(self.period["id"])}
        removed_id = files_before["HDFC_Statement_A.pdf"]
        kept_id = files_before["HDFC_Statement_B.pdf"]

        result = remove_period_file(removed_id)

        self.assertEqual(result["file_id"], removed_id)
        self.assertEqual(result["period_id"], self.period["id"])

        remaining = list_period_files(self.period["id"])
        self.assertEqual([item["original_name"] for item in remaining], ["HDFC_Statement_B.pdf"])
        self.assertEqual(remaining[0]["parse_outcome"], "processed")

        session = get_session()
        try:
            self.assertIsNone(session.get(StoredFile, removed_id))
            self.assertEqual(
                session.query(ExtractedRow).filter(ExtractedRow.file_id == removed_id).count(),
                0,
            )
            self.assertGreater(
                session.query(ExtractedRow).filter(ExtractedRow.file_id == kept_id).count(),
                0,
            )
        finally:
            session.close()

        preview = get_period_preview(self.period["id"], 10)
        preview_ids = {item["file_id"] for item in preview["files"]}
        self.assertNotIn(removed_id, preview_ids)
        self.assertIn(kept_id, preview_ids)

    def test_removing_one_file_drops_its_rows_from_the_rebuilt_pack(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.dump import list_period_files
        from apps.engine.pipeline import get_period_pack
        from apps.engine.removal import remove_period_file

        first, second = self._two_bank_files()
        self._ingest(first, second)

        pack_before = get_period_pack(self.period["id"])
        rows_before = pack_before["row_count"]
        names_before = {item["filename"] for item in pack_before["files"]}
        self.assertEqual(names_before, {"HDFC_Statement_A.pdf", "HDFC_Statement_B.pdf"})

        removed_id = next(
            item["id"]
            for item in list_period_files(self.period["id"])
            if item["original_name"] == "HDFC_Statement_A.pdf"
        )
        remove_period_file(removed_id)

        pack_after = get_period_pack(self.period["id"])
        self.assertEqual({item["filename"] for item in pack_after["files"]}, {"HDFC_Statement_B.pdf"})
        self.assertLess(pack_after["row_count"], rows_before)

        bank_xlsx = next(item["path"] for item in pack_after["outputs"] if item["key"] == "bank")
        workbook = load_workbook(bank_xlsx, read_only=True)
        try:
            text = " ".join(
                str(cell)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
                for cell in row
                if cell is not None
            )
        finally:
            workbook.close()
        self.assertNotIn("HDFC_Statement_A.pdf", text)
        self.assertIn("HDFC_Statement_B.pdf", text)

    def test_removing_the_last_file_leaves_no_pack_behind(self) -> None:
        from apps.engine.db import DataPack, ExtractedRow, get_session
        from apps.engine.dump import list_period_files
        from apps.engine.pipeline import get_period_pack
        from apps.engine.removal import remove_period_file

        only = write_hdfc(self.inbox)
        self._ingest(only)

        pack_before = get_period_pack(self.period["id"])
        self.assertTrue(pack_before["exists"])
        stale_xlsx = [
            Path(item["path"]) for item in pack_before["outputs"] if str(item.get("path", "")).endswith(".xlsx")
        ]

        only_id = list_period_files(self.period["id"])[0]["id"]
        remove_period_file(only_id)

        self.assertEqual(list_period_files(self.period["id"]), [])
        self.assertIsNone(get_period_pack(self.period["id"]))
        session = get_session()
        try:
            self.assertEqual(session.query(DataPack).count(), 0)
            self.assertEqual(session.query(ExtractedRow).count(), 0)
        finally:
            session.close()
        for path in stale_xlsx:
            self.assertFalse(path.exists(), f"{path.name} was left behind")

    def test_removal_is_refused_while_the_period_is_busy(self) -> None:
        from apps.engine.dump import list_period_files, start_job
        from apps.engine.removal import FileRemovalError, remove_period_file

        first, second = self._two_bank_files()
        self._ingest(first, second)
        file_id = list_period_files(self.period["id"])[0]["id"]

        # Simulate a run in flight: an active job exists for the period.
        start_job(self.period["id"])
        with self.assertRaises(FileRemovalError):
            remove_period_file(file_id)

        self.assertEqual(len(list_period_files(self.period["id"])), 2)

    def test_failed_rebuild_restores_the_file_and_reports_no_success(self) -> None:
        from apps.engine.db import StoredFile, get_session
        from apps.engine.dump import list_period_files
        from apps.engine.removal import FileRemovalError, remove_period_file

        first, second = self._two_bank_files()
        self._ingest(first, second)
        removed_id = next(
            item["id"]
            for item in list_period_files(self.period["id"])
            if item["original_name"] == "HDFC_Statement_A.pdf"
        )

        with patch("apps.engine.removal.reparse_period", side_effect=RuntimeError("pack rebuild failed")):
            with self.assertRaises(FileRemovalError):
                remove_period_file(removed_id)

        session = get_session()
        try:
            row = session.get(StoredFile, removed_id)
            self.assertIsNotNone(row)
            self.assertTrue(row.storage_key)
            stored = Path(os.environ["LOCALAPPDATA"]) / "CAUnpacker" / "files" / row.storage_key
            self.assertTrue(stored.exists(), "the stored copy must come back")
        finally:
            session.close()
        self.assertEqual(len(list_period_files(self.period["id"])), 2)

    def test_desktop_api_returns_the_refreshed_desk_for_remaining_files(self) -> None:
        from apps.desktop.app import DesktopApi
        from apps.engine.dump import list_period_files

        first, second = self._two_bank_files()
        self._ingest(first, second)
        removed_id = next(
            item["id"]
            for item in list_period_files(self.period["id"])
            if item["original_name"] == "HDFC_Statement_A.pdf"
        )

        result = DesktopApi().remove_period_file(removed_id)

        self.assertTrue(result["ok"])
        self.assertEqual(result["removed_file_id"], removed_id)
        self.assertEqual(
            [item["original_name"] for item in result["files"]], ["HDFC_Statement_B.pdf"]
        )
        self.assertEqual(
            {item["filename"] for item in (result["pack"] or {}).get("files", [])},
            {"HDFC_Statement_B.pdf"},
        )
        self.assertTrue(result["preview"]["files"])

    def test_desktop_api_never_reports_success_when_removal_fails(self) -> None:
        from apps.desktop.app import DesktopApi
        from apps.engine.dump import list_period_files

        first, second = self._two_bank_files()
        self._ingest(first, second)
        removed_id = list_period_files(self.period["id"])[0]["id"]

        with patch("apps.engine.removal.reparse_period", side_effect=RuntimeError("boom")):
            result = DesktopApi().remove_period_file(removed_id)

        self.assertFalse(result["ok"])
        self.assertTrue(result["error"])
        self.assertNotIn("removed_file_id", result)
        self.assertEqual(len(list_period_files(self.period["id"])), 2)


if __name__ == "__main__":
    unittest.main()
