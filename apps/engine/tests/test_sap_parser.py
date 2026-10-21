from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from apps.engine.tests.fixtures_sap import (
    fixture,
    write_positioned_pdf,
    write_text_pdf,
    write_xlsx,
)


def _write_csv(path: Path, text: str) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


class SapReportTypeTests(unittest.TestCase):
    def test_general_ledger_preserves_formatted_source_values(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_gl_line_items.csv"))
        self.assertEqual(parsed["report_type"], "general_ledger")
        self.assertEqual(parsed["warnings"], [])
        self.assertEqual(len(parsed["rows"]), 2)

        row = parsed["rows"][0]
        self.assertEqual(row["gl_account"], "400000")
        self.assertEqual(row["posting_date"], "2026-07-01")
        self.assertEqual(row["document_number"], "0012345678")
        self.assertEqual(row["debit"], 1250.0)
        self.assertEqual(row["credit"], 0.0)
        # The numeric value is normalised, but the source text is untouched.
        self.assertEqual(row["source_values"]["Debit"], "1,250.00")
        self.assertEqual(row["source_values"]["Document Number"], "0012345678")

    def test_journal_entries_are_grouped_and_balanced(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_journal_entries.csv"))
        self.assertEqual(parsed["report_type"], "journal_entries")
        self.assertEqual(len(parsed["rows"]), 4)
        self.assertEqual(parsed["warnings"], [])
        documents = {row["document_number"] for row in parsed["rows"]}
        self.assertEqual(documents, {"5100000201", "5100000202"})

    def test_trial_balance_keeps_opening_debit_credit_and_closing(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_trial_balance.csv"))
        self.assertEqual(parsed["report_type"], "trial_balance")
        self.assertEqual(len(parsed["rows"]), 6)
        self.assertEqual(parsed["warnings"], [])
        totals = parsed["validation"]["report"]["totals"]
        self.assertEqual(totals["debit"], 65000.0)
        self.assertEqual(totals["credit"], 65000.0)
        cash = next(row for row in parsed["rows"] if row["gl_account"] == "100000")
        self.assertEqual(cash["opening_balance"], 50000.0)
        self.assertEqual(cash["closing_balance"], 70000.0)

    def test_vendor_ledger_uses_vendor_fields(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_vendor_ledger.csv"))
        self.assertEqual(parsed["report_type"], "vendor")
        row = parsed["rows"][0]
        self.assertEqual(row["party_type"], "vendor")
        self.assertEqual(row["vendor_number"], "V-0001")
        self.assertEqual(row["party_code"], "V-0001")
        self.assertEqual(row["supplier_name"], "Northwind Supplies")

    def test_customer_ledger_uses_customer_fields(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_customer_ledger.csv"))
        self.assertEqual(parsed["report_type"], "customer")
        row = parsed["rows"][0]
        self.assertEqual(row["party_type"], "customer")
        self.assertEqual(row["customer_number"], "C-0001")
        self.assertEqual(row["customer_name"], "Bluebird Retail")

    def test_ap_aging_buckets_and_totals(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_ap_aging.csv"))
        self.assertEqual(parsed["report_type"], "accounts_payable")
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["warnings"], [])
        row = parsed["rows"][0]
        self.assertEqual(row["aging_current"], 500.0)
        self.assertEqual(row["aging_1_30"], 400.0)
        self.assertEqual(row["aging_over_120"], 150.0)
        self.assertEqual(row["balance"], 1800.0)
        self.assertEqual(parsed["validation"]["report"]["totals"]["balance"], 3300.0)

    def test_ar_aging_buckets_and_totals(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_ar_aging.csv"))
        self.assertEqual(parsed["report_type"], "accounts_receivable")
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["warnings"], [])
        self.assertEqual(parsed["rows"][0]["aging_91_120"], 200.0)
        self.assertEqual(parsed["validation"]["report"]["totals"]["balance"], 3500.0)

    def test_generic_ledger_is_kept_under_other_with_warning(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        parsed = parse_sap_file(fixture("synthetic_other_ledger.csv"))
        self.assertEqual(parsed["report_type"], "other")
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertTrue(
            any("report type was not identified" in warning for warning in parsed["warnings"])
        )


class SapFormatTests(unittest.TestCase):
    def test_xlsx_title_and_text_document_number(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = write_xlsx(
                Path(tmp) / "synthetic_journal.xlsx",
                {
                    "Journal": [
                        ["SAP Journal Entries - SYNTHETIC TEST DATA"],
                        [
                            "Company Code",
                            "Document Number",
                            "Posting Date",
                            "G/L Account",
                            "Description",
                            "Debit",
                            "Credit",
                        ],
                        [1000, "0012345678", "2026-07-02", "400000", "Travel", 800, 0],
                        [1000, "0012345678", "2026-07-02", "100000", "Payment", 0, 800],
                    ]
                },
            )
            parsed = parse_sap_file(path, "synthetic_journal.xlsx")

        self.assertEqual(parsed["report_type"], "journal_entries")
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["warnings"], [])
        self.assertEqual(parsed["rows"][0]["document_number"], "0012345678")
        self.assertEqual(parsed["rows"][0]["source"], "synthetic_journal.xlsx:Journal")
        self.assertIn("Journal", parsed["validation"])

    def test_positioned_pdf_keeps_debit_and_credit_columns_apart(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = write_positioned_pdf(
                Path(tmp) / "synthetic_gl.pdf",
                title="SAP General Ledger Line Item Report - SYNTHETIC TEST DATA",
                columns=[
                    (40, "Company Code"),
                    (150, "GL Account"),
                    (260, "Posting Date"),
                    (370, "Document Number"),
                    (480, "Debit"),
                    (570, "Credit"),
                ],
                rows=[
                    ["1000", "400000", "2026-07-01", "0012345678", "1,250.00", "0.00"],
                    ["1000", "400000", "2026-07-05", "0012345679", "0.00", "750.50"],
                ],
            )
            parsed = parse_sap_file(path, "synthetic_gl.pdf")

        self.assertEqual(parsed["report_type"], "general_ledger")
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["warnings"], [])
        self.assertEqual(parsed["rows"][0]["debit"], 1250.0)
        self.assertEqual(parsed["rows"][0]["credit"], 0.0)
        self.assertEqual(parsed["rows"][0]["source_values"]["debit"], "1,250.00")
        self.assertEqual(parsed["rows"][0]["source"], "synthetic_gl.pdf#p1")
        self.assertEqual(parsed["rows"][1]["credit"], 750.5)

    def test_text_line_pdf_falls_back_to_line_columns(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = write_text_pdf(
                Path(tmp) / "synthetic_gl_text.pdf",
                [
                    "SAP General Ledger Line Item Report - SYNTHETIC TEST DATA",
                    "Company Code GL Account Posting Date Document Number Debit Credit",
                    "1000 400000 2026-07-01 0012345678 1,250.00 0.00",
                    "1000 400000 2026-07-05 0012345679 0.00 750.50",
                ],
            )
            parsed = parse_sap_file(path, "synthetic_gl_text.pdf")

        self.assertEqual(parsed["report_type"], "general_ledger")
        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["rows"][0]["company_code"], "1000")
        self.assertEqual(parsed["rows"][0]["gl_account"], "400000")
        self.assertEqual(parsed["rows"][0]["document_number"], "0012345678")
        self.assertEqual(parsed["rows"][0]["debit"], 1250.0)
        self.assertEqual(parsed["rows"][1]["credit"], 750.5)

    def test_text_line_pdf_indicator_derives_debit_and_credit(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = write_text_pdf(
                Path(tmp) / "synthetic_gl_indicator.pdf",
                [
                    "SAP General Ledger Indicator - SYNTHETIC TEST DATA",
                    "Company Code Document Number Posting Date Amount Debit/Credit",
                    "1000 D1 2026-07-01 100.00 S",
                    "1000 D2 2026-07-01 50.00 H",
                ],
            )
            parsed = parse_sap_file(path, "synthetic_gl_indicator.pdf")

        self.assertEqual(len(parsed["rows"]), 2)
        self.assertEqual(parsed["rows"][0]["debit"], 100.0)
        self.assertIsNone(parsed["rows"][0]["credit"])
        self.assertIsNone(parsed["rows"][1]["debit"])
        self.assertEqual(parsed["rows"][1]["credit"], 50.0)


class SapValidationTests(unittest.TestCase):
    def _parse_csv(self, text: str, name: str = "synthetic_validation.csv") -> dict:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = _write_csv(Path(tmp) / name, text)
            return parse_sap_file(path, name)

    def test_total_mismatch_warns_without_changing_rows(self) -> None:
        parsed = self._parse_csv(
            "Company Code,GL Account,Posting Date,Document Number,Debit,Credit\n"
            '1000,400000,2026-07-01,D1,"1,250.00",0.00\n'
            'Total,,,,"1,000.00",\n'
        )
        self.assertIn("sap_total_mismatch:debit", parsed["warnings"])
        self.assertEqual(len(parsed["rows"]), 1)
        self.assertEqual(parsed["rows"][0]["debit"], 1250.0)

    def test_debit_credit_total_mismatch_warns(self) -> None:
        parsed = self._parse_csv(
            "Company Code,GL Account,Posting Date,Document Number,Debit,Credit\n"
            '1000,400000,2026-07-01,D1,"1,250.00",0.00\n'
            'Total,,,,"1,250.00","1,000.00"\n'
        )
        self.assertIn("sap_debit_credit_mismatch", parsed["warnings"])
        self.assertNotIn("sap_total_mismatch:debit", parsed["warnings"])

    def test_balance_mismatch_flags_the_row(self) -> None:
        parsed = self._parse_csv(
            "Company Code,GL Account,Posting Date,Document Number,"
            "Opening Balance,Debit,Credit,Closing Balance\n"
            "1000,100000,2026-07-01,D1,100.00,50.00,0.00,200.00\n"
        )
        self.assertIn("sap_balance_mismatch", parsed["warnings"])
        self.assertIn("sap_balance_mismatch", parsed["rows"][0]["flags"])

    def test_unbalanced_journal_document_is_flagged(self) -> None:
        parsed = self._parse_csv(
            "SAP Journal Entries - SYNTHETIC TEST DATA\n"
            "Company Code,Document Number,Posting Date,GL Account,Debit,Credit\n"
            "1000,D1,2026-07-01,400000,100.00,0.00\n"
            "1000,D1,2026-07-01,100000,0.00,90.00\n"
        )
        self.assertIn("sap_document_unbalanced:D1", parsed["warnings"])
        for row in parsed["rows"]:
            self.assertIn("sap_document_unbalanced", row["flags"])

    def test_aging_bucket_mismatch_warns(self) -> None:
        parsed = self._parse_csv(
            "SAP Accounts Payable Aging - SYNTHETIC TEST DATA\n"
            "Company Code,Vendor Number,Vendor Name,Balance,Current,1-30 Days\n"
            "1000,V-1,Acme,100.00,20.00,30.00\n"
        )
        self.assertIn("sap_aging_total_mismatch", parsed["warnings"])
        self.assertIn("sap_aging_total_mismatch", parsed["rows"][0]["flags"])

    def test_trial_balance_without_totals_row_checks_debit_and_credit(self) -> None:
        parsed = self._parse_csv(
            "SAP Trial Balance - SYNTHETIC TEST DATA\n"
            "Company Code,GL Account,Account Name,Debit,Credit\n"
            "1000,100000,Cash,100.00,0.00\n"
            "1000,500000,Revenue,0.00,90.00\n"
        )
        self.assertEqual(parsed["report_type"], "trial_balance")
        self.assertIn("sap_debit_credit_mismatch", parsed["warnings"])

    def test_debit_credit_indicator_derives_side(self) -> None:
        parsed = self._parse_csv(
            "Company Code,Document Number,Posting Date,Amount,Debit/Credit\n"
            "1000,D1,2026-07-01,100.00,S\n"
            "1000,D2,2026-07-01,50.00,H\n"
        )
        first, second = parsed["rows"]
        self.assertEqual(first["debit"], 100.0)
        self.assertIsNone(first["credit"])
        self.assertIsNone(second["debit"])
        self.assertEqual(second["credit"], 50.0)
        self.assertEqual(first["source_values"]["Amount"], "100.00")

    def test_lone_amount_is_not_split_into_debit_or_credit(self) -> None:
        parsed = self._parse_csv(
            "Company Code,Document Number,Posting Date,Amount\n"
            "1000,D1,2026-07-01,100.00\n"
        )
        row = parsed["rows"][0]
        self.assertEqual(row["amount"], 100.0)
        self.assertIsNone(row["debit"])
        self.assertIsNone(row["credit"])


class SapClassificationTests(unittest.TestCase):
    def test_synthetic_csv_reports_are_detected_as_sap(self) -> None:
        from apps.engine.classifier import classify_path

        names = [
            "synthetic_gl_line_items.csv",
            "synthetic_journal_entries.csv",
            "synthetic_trial_balance.csv",
            "synthetic_vendor_ledger.csv",
            "synthetic_customer_ledger.csv",
            "synthetic_ap_aging.csv",
            "synthetic_ar_aging.csv",
        ]
        for name in names:
            with self.subTest(name=name):
                result = classify_path(fixture(name))
                self.assertEqual(result.kind, "sap", result.reason)
                self.assertIn("SAP export structure", result.reason)

    def test_xlsx_and_pdf_structure_is_detected_as_sap(self) -> None:
        from apps.engine.classifier import classify_path

        with tempfile.TemporaryDirectory() as tmp:
            xlsx = write_xlsx(
                Path(tmp) / "synthetic_ledger.xlsx",
                {
                    "Ledger": [
                        ["SAP General Ledger - SYNTHETIC TEST DATA"],
                        [
                            "Company Code",
                            "GL Account",
                            "Posting Date",
                            "Document Number",
                            "Debit",
                            "Credit",
                        ],
                        [1000, "400000", "2026-07-01", "0012345678", 100, 0],
                    ]
                },
            )
            pdf = write_positioned_pdf(
                Path(tmp) / "synthetic_ledger.pdf",
                title="SAP General Ledger - SYNTHETIC TEST DATA",
                columns=[
                    (40, "Company Code"),
                    (150, "GL Account"),
                    (260, "Posting Date"),
                    (370, "Document Number"),
                    (480, "Debit"),
                    (570, "Credit"),
                ],
                rows=[["1000", "400000", "2026-07-01", "0012345678", "100.00", "0.00"]],
            )
            self.assertEqual(classify_path(xlsx).kind, "sap")
            self.assertEqual(classify_path(pdf).kind, "sap")

    def test_bank_and_invoice_text_are_not_reclassified_as_sap(self) -> None:
        from apps.engine.classifier import classify_from_text

        bank = classify_from_text(
            "statement.pdf",
            3,
            "Opening Balance 100.00\n"
            "01/07/2026 Withdrawal 20.00 80.00\n"
            "02/07/2026 Deposit 50.00 130.00\n"
            "Closing Balance 130.00",
        )
        invoice = classify_from_text(
            "bill.pdf",
            1,
            "TAX INVOICE\nInvoice No 88\nGSTIN 27AAPFU0939F1ZV\n"
            "HSN 9983\nTaxable value 10000.00",
        )
        self.assertEqual(bank.kind, "bank")
        self.assertEqual(invoice.kind, "invoice")

    def test_generic_date_description_amount_csv_is_not_sap(self) -> None:
        from apps.engine.classifier import classify_path

        with tempfile.TemporaryDirectory() as tmp:
            path = _write_csv(
                Path(tmp) / "generic.csv",
                "Date,Description,Amount\n2026-07-01,Office supplies,1250.00\n",
            )
            self.assertNotEqual(classify_path(path).kind, "sap")

    def test_zoho_classification_is_unchanged(self) -> None:
        from apps.engine.classifier import classify_path

        with tempfile.TemporaryDirectory() as tmp:
            path = _write_csv(
                Path(tmp) / "zoho_books.csv",
                "Invoice Number,Invoice Date,GST Treatment\nINV-1,2026-07-01,taxable\n",
            )
            self.assertEqual(classify_path(path).kind, "zoho")


class SapUnsupportedLayoutTests(unittest.TestCase):
    def test_legacy_xls_returns_export_guidance(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "legacy.xls"
            path.write_bytes(b"legacy spreadsheet")
            parsed = parse_sap_file(path)
        self.assertEqual(parsed["rows"], [])
        self.assertIn(".xlsx or .csv", parsed["reason"])

    def test_csv_without_financial_header_returns_reason(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = _write_csv(
                Path(tmp) / "notes.csv",
                "Reference,Description\nREF-1,Not a financial report\n",
            )
            parsed = parse_sap_file(path)
        self.assertEqual(parsed["rows"], [])
        self.assertIn("No supported SAP financial table header", parsed["reason"])

    def test_empty_file_returns_reason(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.csv"
            path.write_text("", encoding="utf-8")
            parsed = parse_sap_file(path)
        self.assertEqual(parsed["rows"], [])
        self.assertTrue(parsed.get("reason"))

    def test_garbage_xlsx_returns_error_without_raising(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "broken.xlsx"
            path.write_bytes(b"not-a-spreadsheet")
            parsed = parse_sap_file(path)
        self.assertEqual(parsed["rows"], [])
        self.assertTrue(parsed.get("error"))

    def test_recognised_header_without_transaction_rows_returns_reason(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = _write_csv(
                Path(tmp) / "header_only.csv",
                "Company Code,GL Account,Posting Date,Debit,Credit\n"
                "footer,,,,\n",
            )
            parsed = parse_sap_file(path)
        self.assertEqual(parsed["rows"], [])
        self.assertIn("no usable transaction rows", parsed["reason"])

    def test_pdf_without_table_returns_reason(self) -> None:
        from apps.engine.parsers.sap import parse_sap_file

        with tempfile.TemporaryDirectory() as tmp:
            path = write_text_pdf(
                Path(tmp) / "note.pdf",
                ["This is a plain note.", "No financial columns are present."],
            )
            parsed = parse_sap_file(path)
        self.assertEqual(parsed["rows"], [])
        self.assertIn("No supported SAP table", parsed["reason"])

    def test_sap_outcome_records_identity_and_no_rows_review(self) -> None:
        from apps.engine.outcomes import FileOutcome, evaluate_file_outcome

        result = evaluate_file_outcome(
            kind="sap",
            rows=[],
            parser_metadata={"reason": "No supported SAP financial table header was found."},
        )
        self.assertEqual(result.outcome, FileOutcome.NEEDS_REVIEW)
        self.assertEqual(result.reason_code, "no_rows")
        self.assertEqual(result.parser_id, "sap")
        self.assertTrue(result.parser_version)


class SapPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._tmp.cleanup()

    def test_sap_csv_writes_separate_sap_workbook(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.clients import create_client
        from apps.engine.dump import ingest_paths, start_job
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period
        from apps.engine.pipeline import get_period_pack

        save_firm("Test firm")
        client = create_client("Acme")
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(fixture("synthetic_gl_line_items.csv"))])

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        by_key = {item["key"]: item for item in pack["outputs"]}
        self.assertIn("sap", by_key)
        self.assertNotIn("purchase", by_key)
        self.assertNotIn("sales", by_key)
        self.assertNotIn("books", by_key)

        sap_output = by_key["sap"]
        self.assertEqual(sap_output["label"], "SAP_Reports_Extracted.xlsx")
        path = Path(sap_output["path"])
        self.assertEqual(path.name, "SAP_Reports_Extracted.xlsx")
        self.assertTrue(path.is_file())
        self.assertEqual(sap_output["rows"], 2)

        book = load_workbook(path, data_only=True)
        try:
            sheet = book.active
            headers = [cell for cell in next(sheet.iter_rows(values_only=True))]
        finally:
            book.close()
        self.assertIn("G/L account", headers)
        self.assertIn("Debit", headers)
        self.assertIn("Credit", headers)
        self.assertEqual(pack["row_count"], 2)

    def test_unsupported_csv_overridden_to_sap_needs_review(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.clients import create_client
        from apps.engine.dump import (
            ingest_paths,
            list_period_files,
            override_kind,
            reparse_period,
            start_job,
        )
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period
        from apps.engine.pipeline import get_period_pack

        inbox = Path(self._tmp.name) / "inbox"
        inbox.mkdir()
        notes = inbox / "notes.csv"
        notes.write_text("Reference,Description\nREF-1,Not a financial report\n", encoding="utf-8")

        save_firm("Test firm")
        client = create_client("Acme")
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(notes)])

        stored = next(
            item
            for item in list_period_files(period["id"])
            if item["original_name"] == "notes.csv"
        )
        override_kind(stored["id"], "sap")
        reparsed = reparse_period(period["id"])
        self.assertNotEqual(reparsed.get("status"), "failed", reparsed)

        updated = next(
            item
            for item in list_period_files(period["id"])
            if item["original_name"] == "notes.csv"
        )
        self.assertEqual(updated["kind"], "sap")
        self.assertEqual(updated["parse_outcome"], "needs_review")
        self.assertEqual(updated["parse_reason_code"], "no_rows")
        self.assertIn("No supported SAP financial table header", updated["classify_reason"])
        self.assertTrue(updated["needs_review"])

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        by_key = {item["key"]: item for item in pack["outputs"]}
        self.assertNotIn("sap", by_key)
        self.assertIn("needs_review", by_key)

        review_path = Path(by_key["needs_review"]["path"])
        book = load_workbook(review_path, data_only=True)
        try:
            text = " ".join(
                str(cell)
                for row in book.active.iter_rows(values_only=True)
                for cell in row
                if cell is not None
            )
        finally:
            book.close()
        self.assertIn("No supported SAP financial table header", text)


if __name__ == "__main__":
    unittest.main()
