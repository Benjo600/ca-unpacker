from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from apps.engine.file_review import (
    ATTACHED,
    ESSENTIAL,
    MISSING,
    RECOMMENDED,
    VERIFIED,
    DocNeed,
    _needs_for,
    evaluate_file_review,
)
from apps.engine.pipeline import get_period_pack
from apps.engine.clients import create_client
from apps.engine.dump import ingest_paths, start_job
from apps.engine.firm import save_firm
from apps.engine.periods import create_period


def _profile(profession_key: str = "doctor", gst_registered: bool | None = False) -> dict:
    return {
        "rules": {},
        "profession_key": profession_key,
        "profession_label": "Doctor / clinic",
        "summary": "Doctor / clinic",
        "entity_type_label": "Proprietorship",
        "gst_registered": gst_registered,
        "gstin": "",
    }


def _file(
    *,
    name: str,
    kind: str,
    outcome: str,
    reason: str = "",
) -> SimpleNamespace:
    return SimpleNamespace(
        original_name=name,
        detected_kind=kind,
        override_kind=None,
        parse_outcome=outcome,
        parse_reason_message=reason,
    )


def _bank_row(date: str, debit: float | None = None, credit: float | None = None, text: str = "") -> dict:
    row = {"date": date, "description": text}
    if debit is not None:
        row["debit"] = debit
    if credit is not None:
        row["credit"] = credit
    return row


class ChecklistCatalogTests(unittest.TestCase):
    def test_four_profession_packs_plus_generic_exist(self) -> None:
        for key in ("doctor", "trader_wholesale", "trader_retail", "consultant", "contractor", "astronaut_unlisted"):
            needs = _needs_for(key)
            self.assertTrue(needs, key)
            for need in needs:
                self.assertIn(need.grade, (ESSENTIAL, RECOMMENDED))
                self.assertTrue(need.why, key)
                for kind in need.read_kinds:
                    self.assertIn(kind, ("bank", "invoice", "gstr_1", "gstr_2b", "gstr_3b", "tally", "zoho"))

    def test_needed_keys_are_unique(self) -> None:
        for key in ("doctor", "trader_wholesale", "trader_retail", "consultant", "contractor", "anything"):
            keys = [need.key for need in _needs_for(key)]
            self.assertEqual(len(keys), len(set(keys)), key)


class ChecklistOutcomeTests(unittest.TestCase):
    def test_read_documents_verify_the_need(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=[
                _file(name="HDFC_Jul.pdf", kind="bank", outcome="processed"),
                _file(name="lab bill.pdf", kind="invoice", outcome="processed"),
            ],
            bank_rows=[_bank_row("2026-07-03", debit=5000, text="NEFT")],
            purchase_rows=[
                {
                    "supplier_name": "Lab One",
                    "invoice_value": 4000,
                    "invoice_date": "2026-07-10",
                }
            ],
        )
        statuses = {item["key"]: item["status"] for item in review["checklist"]}
        self.assertEqual(statuses["bank_statement"], VERIFIED)
        self.assertEqual(statuses["expense_bills"], VERIFIED)

    def test_failed_bank_file_is_attached_not_read(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=[
                _file(
                    name="HDFC_Jul.pdf",
                    kind="bank",
                    outcome="failed",
                    reason="could not parse: bad table",
                ),
            ],
        )
        item = next(x for x in review["checklist"] if x["key"] == "bank_statement")
        self.assertEqual(item["status"], ATTACHED)
        self.assertIn("could not read", item["evidence"])
        self.assertGreaterEqual(review["counts"]["attached_not_read"], 1)

    def test_missing_essential_packages_readiness_gaps(self) -> None:
        review = evaluate_file_review(profile=_profile(), files=[])
        item = next(x for x in review["checklist"] if x["key"] == "bank_statement")
        self.assertEqual(item["status"], MISSING)
        self.assertEqual(review["readiness"], "gaps")
        self.assertGreaterEqual(review["counts"]["essential_missing"], 1)

    def test_attach_only_document_is_found_by_file_name(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=[_file(name="Form16A_Jul2026.pdf", kind="unknown", outcome="unclassified")],
        )
        item = next(x for x in review["checklist"] if x["key"] == "tds_certificates")
        self.assertEqual(item["status"], ATTACHED)
        self.assertIn("file name only", item["evidence"])

    def test_sales_rows_verify_the_sales_record(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            sales_rows=[
                {"invoice_value": 9000, "invoice_date": "2026-07-05"},
            ],
        )
        item = next(x for x in review["checklist"] if x["key"] == "sales_record")
        self.assertEqual(item["status"], VERIFIED)


class GstSwitchTests(unittest.TestCase):
    def test_registered_clients_get_the_portal_downloads(self) -> None:
        review = evaluate_file_review(profile=_profile(gst_registered=True), files=[])
        keys = {item["key"] for item in review["checklist"]}
        self.assertIn("gstr_2b_download", keys)
        self.assertIn("gstr_1_download", keys)
        self.assertIn("gstr_3b_download", keys)

    def test_unregistered_clients_do_not(self) -> None:
        review = evaluate_file_review(profile=_profile(gst_registered=False), files=[])
        keys = {item["key"] for item in review["checklist"]}
        self.assertNotIn("gstr_2b_download", keys)
        self.assertNotIn("gstr_3b_download", keys)


class EscalationTests(unittest.TestCase):
    def test_rent_treatment_flag_upgrades_the_rent_agreement(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=[],
            treatment_flags=[
                {
                    "code": "tds_194i",
                    "severity": "review",
                    "title": "Possible 194I payments",
                }
            ],
        )
        item = next(x for x in review["checklist"] if x["key"] == "rent_agreement")
        self.assertEqual(item["grade"], ESSENTIAL)
        self.assertTrue(item["escalated"])

    def test_salary_payment_rows_upgrade_the_salary_register(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=[],
            bank_rows=[
                _bank_row("2026-07-28", debit=45000, text="SALARY JULY STAFF"),
            ],
        )
        item = next(x for x in review["checklist"] if x["key"] == "salary_register")
        self.assertEqual(item["grade"], ESSENTIAL)
        self.assertTrue(item["escalated"])


class CrossCheckTests(unittest.TestCase):
    def test_matching_purchases_and_2b_raise_nothing(self) -> None:
        purchase = [
            {"supplier_name": "Acme", "invoice_value": 1180, "taxable_value": 1000},
        ]
        portal = [
            {"gstin": "29ABCDE1234F1Z5", "trade_name": "Acme", "invoice_value": 1180, "taxable": 1000},
        ]
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            purchase_rows=purchase,
            gstr_2b_rows=portal,
        )
        codes = {flag["code"] for flag in review["checks"]}
        self.assertNotIn("review_purchases_vs_2b", codes)

    def test_purchases_below_2b_raise_a_finding(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            purchase_rows=[
                {"supplier_name": "Acme", "invoice_value": 1180, "taxable_value": 1000},
            ],
            gstr_2b_rows=[
                {"gstin": "29ABCDE1234F1Z5", "trade_name": "Acme", "invoice_value": 11800, "taxable": 10000},
            ],
        )
        flag = next(x for x in review["checks"] if x["code"] == "review_purchases_vs_2b")
        self.assertLess(flag["observed"]["difference"], 0)
        self.assertIn(flag["severity"], ("review", "check"))

    def test_books_sales_versus_gstr1(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            sales_rows=[{"invoice_value": 500000}],
            gstr_1_rows=[{"invoice_value": 250000}],
        )
        codes = {flag["code"] for flag in review["checks"]}
        self.assertIn("review_sales_vs_gstr1", codes)

    def test_supplier_gstin_conflict(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            purchase_rows=[
                {"supplier_name": "Acme Traders", "supplier_gstin": "29AAAAA1111A1Z5", "invoice_value": 1180},
            ],
            gstr_2b_rows=[
                {"trade_name": "Acme Traders", "gstin": "29BBBBB2222B1Z4", "invoice_value": 1180},
            ],
        )
        codes = {flag["code"] for flag in review["checks"]}
        self.assertIn("check_supplier_gstin_conflict", codes)

    def test_duplicate_bill_number_in_purchase_register(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            purchase_rows=[
                {"supplier_name": "Acme", "invoice_number": "101", "invoice_value": 1180},
                {"supplier_name": "Acme", "invoice_number": "101", "invoice_value": 1180},
            ],
            gstr_2b_rows=[
                {"gstin": "29ABCDE1234F1Z5", "invoice_number": "101", "invoice_value": 1180},
            ],
        )
        flag = next(x for x in review["checks"] if x["code"] == "check_duplicate_invoice")
        self.assertEqual(flag["severity"], "review")
        self.assertEqual(flag["observed"]["count"], 2)

    def test_statement_gap_between_rows(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            bank_rows=[
                _bank_row("2026-07-31", debit=100, text="NEFT"),
                _bank_row("2026-08-30", credit=100, text="DEPOSIT"),
            ],
        )
        codes = {flag["code"] for flag in review["checks"]}
        self.assertIn("check_statement_gap", codes)

    def test_cash_payment_above_limit(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            bank_rows=[
                _bank_row("2026-07-14", debit=25000, text="ATM CASH WITHDRAWAL"),
            ],
        )
        codes = {flag["code"] for flag in review["checks"]}
        self.assertIn("check_cash_payments", codes)

    def test_normal_bank_credit_near_sales_raises_nothing(self) -> None:
        review = evaluate_file_review(
            profile=_profile(gst_registered=True),
            sales_rows=[{"invoice_value": 100000}],
            bank_rows=[
                _bank_row("2026-07-05", credit=50000, text="NEFT CUSTOMER"),
                _bank_row("2026-07-20", credit=60000, text="UPI CUSTOMER"),
            ],
        )
        codes = {flag["code"] for flag in review["checks"]}
        self.assertNotIn("check_sales_vs_bank", codes)


class ReadinessTests(unittest.TestCase):
    def _full_doctor_files(self) -> list:
        return [
            _file(name="HDFC.pdf", kind="bank", outcome="processed"),
            _file(name="bills.pdf", kind="invoice", outcome="processed"),
            _file(name="Form 16A_Jul.pdf", kind="unknown", outcome="unclassified"),
            _file(name="Rent Agreement.pdf", kind="unknown", outcome="unclassified"),
            _file(name="Salary Register.xlsx", kind="unknown", outcome="unclassified"),
            _file(name="Tally Export.xml", kind="tally", outcome="processed"),
        ]

    def test_ready_when_everything_read_and_no_checks(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=self._full_doctor_files(),
            bank_rows=[_bank_row("2026-07-01", debit=10, text="NEFT")],
            purchase_rows=[{"supplier_name": "Acme", "invoice_value": 100}],
            sales_rows=[{"invoice_value": 200, "invoice_date": "2026-07-05"}],
            books_rows=[{"voucher_type": "Purchase", "amount": 100}],
        )
        codes = [flag["code"] for flag in review["checks"]]
        self.assertEqual(codes, [])
        self.assertEqual(review["readiness"], "ready")

    def test_check_when_local_checks_only(self) -> None:
        review = evaluate_file_review(
            profile=_profile(),
            files=self._full_doctor_files(),
            bank_rows=[
                _bank_row("2026-07-14", debit=25000, text="CASH ATM"),
                _bank_row("2026-08-01", credit=100, text="DEP"),
            ],
            purchase_rows=[{"supplier_name": "Acme", "invoice_value": 100}],
            sales_rows=[{"invoice_value": 200, "invoice_date": "2026-07-05"}],
            books_rows=[{"voucher_type": "Purchase", "amount": 100}],
        )
        self.assertGreater(review["counts"]["check"], 0)
        self.assertEqual(review["readiness"], "check")


class EndToEndFileReviewTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()
        save_firm("Test firm")

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._tmp.cleanup()

    def _write_statement(self) -> Path:
        from make_test_dump import pdf_with_text

        inbox = Path(self._tmp.name) / "inbox"
        inbox.mkdir()
        statement = inbox / "HDFC_Statement.pdf"
        statement.write_bytes(
            pdf_with_text(
                [
                    "HDFC Bank Account Statement",
                    "Account Number 50100123456789",
                    "Opening Balance 500000.00",
                    "01/07/2026 NEFT Withdrawal PROFESSIONAL FEES DR SHARMA 75000.00 425000.00",
                    "Closing Balance 425000.00",
                ]
            )
        )
        return statement

    def test_pack_carries_the_file_review(self) -> None:
        client = create_client(
            "Dr Mehta",
            "",
            {
                "profession_key": "doctor",
                "entity_type": "proprietorship",
                "gst_registered": False,
            },
        )
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(self._write_statement())])

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        review = pack.get("file_review")
        self.assertIsNotNone(review)
        self.assertIn("readiness", review)
        statuses = {item["key"]: item["status"] for item in review["checklist"]}
        self.assertEqual(statuses["bank_statement"], VERIFIED)
        self.assertEqual(statuses["expense_bills"], MISSING)

        outputs = {item["key"]: item for item in pack["outputs"]}
        self.assertIn("file_review", outputs)
        workbook_path = Path(outputs["file_review"]["path"])
        self.assertTrue(workbook_path.is_file(), workbook_path)
        book = load_workbook(workbook_path)
        self.assertEqual(set(book.sheetnames), {"Summary", "Checklist", "Cross-checks"})

    def test_no_file_review_without_a_profile(self) -> None:
        client = create_client("Acme")
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(self._write_statement())])

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        self.assertIsNone(pack.get("file_review"))


if __name__ == "__main__":
    unittest.main()
