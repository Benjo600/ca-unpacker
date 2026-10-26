from __future__ import annotations

import os
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from apps.engine.clients import create_client, get_client, list_clients, update_client_profile
from apps.engine.dump import ingest_paths, start_job
from apps.engine.firm import save_firm
from apps.engine.periods import create_period
from apps.engine.pipeline import get_period_pack


class ClientProfileStorageTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()
        save_firm("Test firm")

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._tmp.cleanup()

    def test_create_client_saves_the_profile(self) -> None:
        client = create_client(
            "Dr Mehta",
            "",
            {
                "profession_key": "doctor",
                "entity_type": "proprietorship",
                "gst_registered": False,
            },
        )
        self.assertEqual(client["profile"]["profession_key"], "doctor")
        self.assertEqual(client["profile"]["entity_type"], "proprietorship")
        self.assertIs(client["profile"]["gst_registered"], False)
        self.assertTrue(client["profile_is_set"])
        self.assertIn("Doctor / clinic", client["profile_summary"])
        self.assertIn("Not GST registered", client["profile_summary"])

        reloaded = get_client(client["id"])
        self.assertEqual(reloaded["profile"], client["profile"])
        self.assertEqual(len(list_clients()), 1)

    def test_create_client_without_a_profile_still_works(self) -> None:
        client = create_client("Acme")
        self.assertFalse(client["profile_is_set"])
        self.assertEqual(client["profile_summary"], "No profile set")
        self.assertEqual(client["profile"]["profession_key"], "")

    def test_update_replaces_and_can_clear_the_profile(self) -> None:
        client = create_client("Acme")
        updated = update_client_profile(
            client["id"],
            {
                "profession_key": "trader_wholesale",
                "profession_label": "",
                "entity_type": "partnership",
                "gst_registered": True,
                "books": "Tally",
                "banks": "HDFC, ICICI",
            },
        )
        self.assertEqual(updated["profile"]["profession_key"], "trader_wholesale")
        self.assertEqual(updated["profile"]["books"], "Tally")
        self.assertEqual(updated["profile"]["banks"], ["HDFC", "ICICI"])

        cleared = update_client_profile(client["id"], {})
        self.assertFalse(cleared["profile_is_set"])
        self.assertEqual(cleared["profile"]["profession_key"], "")
        self.assertEqual(cleared["profile"]["banks"], [])

    def test_invalid_profile_does_not_create_the_client(self) -> None:
        with self.assertRaisesRegex(ValueError, "not in the list"):
            create_client("Acme", "", {"profession_key": "astronaut"})
        self.assertEqual(list_clients(), [])

    def test_custom_profession_needs_a_label(self) -> None:
        with self.assertRaisesRegex(ValueError, "Describe the profession"):
            create_client("Acme", "", {"profession_key": "custom"})

    def test_unknown_client_update_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "not found"):
            update_client_profile(9999, {"profession_key": "doctor"})


class PeriodPackTreatmentTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()
        save_firm("Test firm")

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        self._tmp.cleanup()

    def _write_statement(self) -> Path:
        from make_test_dump import pdf_with_text

        inbox = Path(self._tmp.name) / "inbox"
        inbox.mkdir()
        statement = inbox / "HDFC_Statement.pdf"
        statement.write_bytes(
            pdf_with_text(
                [
                    "HDFC Bank Account Statement",
                    "Account Number 50100123456789",
                    "Opening Balance 500000.00",
                    "01/07/2026 NEFT Withdrawal PROFESSIONAL FEES DR SHARMA 75000.00 425000.00",
                    "Closing Balance 425000.00",
                ]
            )
        )
        return statement

    def test_pack_carries_treatment_notes_for_a_profiled_client(self) -> None:
        client = create_client(
            "Dr Mehta",
            "",
            {
                "profession_key": "doctor",
                "entity_type": "proprietorship",
                "gst_registered": False,
            },
        )
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(self._write_statement())])

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        treatment = pack.get("treatment")
        self.assertIsNotNone(treatment)
        self.assertIn("Doctor / clinic", treatment["profile"]["summary"])
        self.assertTrue(treatment["disclaimer"])
        self.assertIn(
            "tds_194j",
            {flag["code"] for flag in treatment["flags"]},
            treatment["flags"],
        )

        outputs = {item["key"]: item for item in pack["outputs"]}
        self.assertIn("treatment", outputs)
        workbook_path = Path(outputs["treatment"]["path"])
        self.assertTrue(workbook_path.is_file(), workbook_path)
        book = load_workbook(workbook_path)
        self.assertIn("Treatment notes", book.sheetnames)
        self.assertIn("Expense heads", book.sheetnames)
        self.assertIn("Profile", book.sheetnames)

    def test_pack_has_no_treatment_sheet_without_a_profile(self) -> None:
        client = create_client("Acme")
        period = create_period(client["id"], "Jul 2026")
        job = start_job(period["id"])
        ingest_paths(job["id"], [str(self._write_statement())])

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        self.assertIsNone(pack.get("treatment"))
        keys = {item["key"] for item in pack["outputs"]}
        self.assertNotIn("treatment", keys)


if __name__ == "__main__":
    unittest.main()
