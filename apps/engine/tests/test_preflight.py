from __future__ import annotations

import json
import os
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from apps.engine.preflight import (
    BLOCKING,
    WARNING,
    PreflightIssue,
    PreflightReport,
    PreflightSample,
    load_report,
    run_preflight,
    save_report,
)
from apps.engine.validators.gstin import gstin_flags

VALID_GSTIN = "27AAPFU0939F1ZV"


def _bad_checksum_gstin() -> str:
    """Same 15-character shape as a real GSTIN, wrong check digit."""
    candidate = VALID_GSTIN[:14] + ("W" if VALID_GSTIN[14] != "W" else "X")
    assert gstin_flags(candidate) == ["gstin_checksum"], candidate
    return candidate


def _row(**overrides) -> dict:
    row = {
        "source": "Batch_Jul.xlsx",
        "invoice_number": "INV-1",
        "invoice_date": "05-07-2026",
        "taxable_value": 100.0,
        "tax": 18.0,
        "invoice_value": 118.0,
        "hsn": "0401",
        "supplier_gstin": VALID_GSTIN,
    }
    row.update(overrides)
    return row


def _codes(report: PreflightReport) -> dict[str, PreflightIssue]:
    return {issue.code: issue for issue in report.issues}


class PreflightRuleTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.dest = Path(self._tmp.name) / "pack"

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_clean_material_has_no_findings(self) -> None:
        report = run_preflight(dest_dir=self.dest, purchase_rows=[_row()])

        self.assertTrue(report.ok)
        self.assertFalse(report.blocked)
        self.assertEqual(report.issues, ())
        self.assertEqual(report.counts, {"blocking": 0, "warnings": 0})

    def test_blocking_when_the_period_has_no_usable_rows(self) -> None:
        report = run_preflight(dest_dir=self.dest)

        self.assertFalse(report.ok)
        self.assertTrue(report.blocked)
        self.assertEqual([issue.code for issue in report.blocking], ["no_usable_rows"])
        self.assertEqual(report.blocking[0].level, BLOCKING)
        self.assertEqual(report.warnings, [])

    def test_blocking_when_the_destination_folder_cannot_be_created(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            destination_error="[WinError 5] Access is denied",
            purchase_rows=[_row()],
        )

        self.assertTrue(report.blocked)
        issue = report.blocking[0]
        self.assertEqual(issue.code, "destination_unwritable")
        self.assertIn("Access is denied", issue.samples[0].detail)

    def test_blocking_when_the_destination_probe_cannot_be_written(self) -> None:
        blocker = Path(self._tmp.name) / "a-file"
        blocker.write_text("not a folder", encoding="utf-8")

        report = run_preflight(dest_dir=blocker / "pack", purchase_rows=[_row()])

        self.assertTrue(report.blocked)
        self.assertEqual([issue.code for issue in report.blocking], ["destination_unwritable"])

    def test_gstin_format_warning_counts_rows_and_keeps_samples(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            purchase_rows=[_row(), _row(supplier_gstin="27ABC")],
        )

        issue = _codes(report)["gstin_format"]
        self.assertEqual(issue.level, WARNING)
        self.assertEqual(issue.count, 1)
        self.assertIn("1 row", issue.message)
        self.assertEqual(issue.samples[0].file, "Batch_Jul.xlsx")
        self.assertEqual(issue.samples[0].row, 2)

    def test_gstin_checksum_warning_is_separate_from_format(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            sales_rows=[_row(supplier_gstin=_bad_checksum_gstin())],
        )

        codes = _codes(report)
        self.assertIn("gstin_checksum", codes)
        self.assertNotIn("gstin_format", codes)
        self.assertEqual(codes["gstin_checksum"].count, 1)

    def test_missing_hsn_warning_counts_invoices_and_lines(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            purchase_rows=[_row(hsn=None), _row(hsn="  ")],
            line_item_rows=[{"source": "invoice.pdf", "hsn": "", "taxable": 100.0}],
        )

        issue = _codes(report)["hsn_missing"]
        self.assertEqual(issue.count, 3)
        self.assertIn("3 rows", issue.message)
        self.assertEqual(len(issue.samples), 3)

    def test_missing_date_warning_ignores_summary_rows(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            purchase_rows=[_row(invoice_date=None)],
            gstr_1=[{"gstin": VALID_GSTIN, "document_type": "HSN", "taxable": 500.0}],
        )

        issue = _codes(report)["missing_date"]
        self.assertEqual(issue.count, 1)
        self.assertEqual(issue.samples[0].row, 1)

    def test_invoice_math_warning_allows_one_rupee_rounding(self) -> None:
        rounding_ok = run_preflight(
            dest_dir=self.dest,
            purchase_rows=[_row(invoice_value=118.5)],
        )
        self.assertNotIn("invoice_math", _codes(rounding_ok))

        wrong = run_preflight(dest_dir=self.dest, purchase_rows=[_row(invoice_value=150.0)])
        issue = _codes(wrong)["invoice_math"]
        self.assertEqual(issue.count, 1)
        self.assertIn("150", issue.samples[0].detail)

    def test_invoice_math_uses_gst_split_columns_when_there_is_no_tax_total(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            gstr_2b=[
                {
                    "gstin": VALID_GSTIN,
                    "invoice_number": "2B-1",
                    "invoice_date": "05-07-2026",
                    "invoice_value": 200.0,
                    "taxable": 100.0,
                    "igst": 18.0,
                    "cgst": 0.0,
                    "sgst": 0.0,
                }
            ],
        )

        issue = _codes(report)["invoice_math"]
        self.assertEqual(issue.count, 1)

    def test_duplicate_invoice_numbers_are_counted_per_register(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            purchase_rows=[_row(invoice_number="PUR-1"), _row(invoice_number="PUR-1")],
            sales_rows=[_row(invoice_number="PUR-1")],
        )

        issue = _codes(report)["duplicate_invoice"]
        self.assertEqual(issue.count, 1)
        self.assertEqual(issue.samples[0].row, 2)
        self.assertIn("appears 2 times", issue.samples[0].detail)

    def test_unparsed_files_are_reported_but_clean_files_are_not(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            files=[
                {
                    "original_name": "broken.pdf",
                    "parse_outcome": "needs_review",
                    "parse_reason_code": "no_rows",
                    "parse_reason_message": "The file was recognized, but no rows could be extracted.",
                },
                {
                    "original_name": "HDFC.pdf",
                    "kind": "bank",
                    "parse_outcome": "processed",
                    "parse_reason_code": "rows_extracted",
                },
            ],
            bank_files=[
                {
                    "rows": [{"debit": 1.0}],
                    "check": {"status": "match", "row_count": 1},
                    "meta": {"filename": "HDFC.pdf"},
                }
            ],
        )

        issue = _codes(report)["file_unparsed"]
        self.assertEqual(issue.count, 1)
        self.assertEqual([sample.file for sample in issue.samples], ["broken.pdf"])

    def test_bank_statement_that_did_not_reconcile_is_a_warning(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            bank_files=[
                {
                    "rows": [{"debit": 1.0}, {"debit": 2.0}],
                    "check": {
                        "status": "mismatch",
                        "row_count": 2,
                        "broken_at_row": 2,
                        "flags": ["balance_mismatch"],
                    },
                    "meta": {"filename": "HDFC_Jul2026.pdf"},
                },
                {
                    "rows": [{"debit": 3.0}],
                    "check": {"status": "match", "row_count": 1},
                    "meta": {"filename": "SBI_Jul2026.pdf"},
                },
            ],
        )

        issue = _codes(report)["bank_unreconciled"]
        self.assertEqual(issue.count, 1)
        self.assertEqual(issue.samples[0].file, "HDFC_Jul2026.pdf")
        self.assertEqual(issue.samples[0].row, 2)

    def test_report_round_trips_through_json(self) -> None:
        report = run_preflight(
            dest_dir=self.dest,
            purchase_rows=[_row(hsn=None)],
            files=[{"original_name": "notes.bin", "parse_outcome": "unclassified"}],
        )
        payload = json.loads(json.dumps(report.to_dict()))

        self.assertFalse(payload["blocked"])
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["counts"]["warnings"], 2)

        restored = PreflightReport.from_dict(payload)
        assert restored is not None
        self.assertEqual(
            [(issue.code, issue.level, issue.count) for issue in restored.issues],
            [(issue.code, issue.level, issue.count) for issue in report.issues],
        )
        self.assertIsNone(PreflightReport.from_dict(None))

    def test_warning_samples_stop_at_the_sample_limit_but_the_count_does_not(self) -> None:
        rows = [_row(hsn=None) for _ in range(9)]

        issue = _codes(run_preflight(dest_dir=self.dest, purchase_rows=rows))["hsn_missing"]

        self.assertEqual(issue.count, 9)
        self.assertEqual(len(issue.samples), 5)

    def test_saved_report_is_reloadable_from_the_library(self) -> None:
        from apps.engine.db import reset_engine

        with tempfile.TemporaryDirectory() as library:
            previous = os.environ.get("LOCALAPPDATA")
            os.environ["LOCALAPPDATA"] = library
            try:
                reset_engine()
                report = run_preflight(dest_dir=self.dest, purchase_rows=[_row()])
                path = save_report(4321, report)

                self.assertTrue(path.is_file())
                loaded = load_report(4321)
                assert loaded is not None
                self.assertTrue(loaded.ok)
                self.assertEqual(loaded.issues, ())
            finally:
                reset_engine()
                if previous is None:
                    os.environ.pop("LOCALAPPDATA", None)
                else:
                    os.environ["LOCALAPPDATA"] = previous


class PreflightWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory()
        self.path = Path(self._tmp.name) / "Preflight_Report.xlsx"

    def tearDown(self) -> None:
        self._tmp.cleanup()

    def test_blocking_rows_are_painted_red_and_one_row_per_issue_is_written(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.preflight_xlsx import PREFLIGHT_HEADERS, write_preflight_report
        from apps.engine.pack.style import RED_BG

        report = PreflightReport(
            issues=(
                PreflightIssue(
                    level=BLOCKING,
                    code="destination_unwritable",
                    message="The output folder cannot be written to.",
                    count=1,
                    samples=(PreflightSample(file="C:/packs/Acme"),),
                ),
                PreflightIssue(
                    level=WARNING,
                    code="hsn_missing",
                    message="Missing HSN code on 2 rows.",
                    count=2,
                    samples=(PreflightSample(file="Batch.xlsx", row=4),),
                ),
            )
        )

        write_preflight_report(self.path, report)
        book = load_workbook(self.path)
        sheet = book["Preflight"]

        self.assertEqual([cell.value for cell in sheet[1]], PREFLIGHT_HEADERS)
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet.cell(2, 1).value, "BLOCKING")
        self.assertEqual(sheet.cell(3, 1).value, "WARNING")
        self.assertTrue(str(sheet.cell(2, 1).fill.fgColor.rgb).endswith(RED_BG))
        self.assertFalse(str(sheet.cell(3, 1).fill.fgColor.rgb).endswith(RED_BG))
        self.assertIn("row 4", str(sheet.cell(3, 5).value))

    def test_empty_report_writes_a_single_note_row(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pack.preflight_xlsx import write_preflight_report

        write_preflight_report(self.path, PreflightReport())

        sheet = load_workbook(self.path)["Preflight"]
        self.assertEqual(sheet.max_row, 2)
        self.assertIn("nothing to report", str(sheet.cell(2, 3).value))


TALLY_WITH_ISSUES = """<?xml version="1.0" encoding="utf-8"?>
<ENVELOPE>
  <TALLYMESSAGE>
    <VOUCHER VCHTYPE="Purchase">
      <VOUCHERNUMBER>1</VOUCHERNUMBER>
      <REFERENCE>PUR-1</REFERENCE>
      <PARTYLEDGERNAME>Alpha Traders</PARTYLEDGERNAME>
      <PARTYGSTIN>27AAPFU0939F1ZV</PARTYGSTIN>
      <DATE>20260705</DATE>
      <AMOUNT>1180.00</AMOUNT>
      <TAXABLEVALUE>1000.00</TAXABLEVALUE>
      <TAXAMOUNT>180.00</TAXAMOUNT>
    </VOUCHER>
    <VOUCHER VCHTYPE="Purchase">
      <VOUCHERNUMBER>2</VOUCHERNUMBER>
      <REFERENCE>PUR-1</REFERENCE>
      <PARTYLEDGERNAME>Alpha Traders</PARTYLEDGERNAME>
      <PARTYGSTIN>27AAPFU0939F1ZV</PARTYGSTIN>
      <AMOUNT>1500.00</AMOUNT>
      <TAXABLEVALUE>1000.00</TAXABLEVALUE>
      <TAXAMOUNT>180.00</TAXAMOUNT>
    </VOUCHER>
  </TALLYMESSAGE>
</ENVELOPE>
"""


class PipelinePreflightTests(unittest.TestCase):
    def setUp(self) -> None:
        from apps.engine.db import reset_engine
        from apps.engine.pdf_passwords import clear_all_passwords

        self._tmp = tempfile.TemporaryDirectory()
        self.previous_localappdata = os.environ.get("LOCALAPPDATA")
        os.environ["LOCALAPPDATA"] = self._tmp.name
        reset_engine()
        clear_all_passwords()

        from apps.engine.clients import create_client
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period

        save_firm("Preflight firm")
        client = create_client("Acme")
        self.period = create_period(client["id"], "Jul 2026")
        self.inbox = Path(self._tmp.name) / "inbox"
        self.inbox.mkdir()

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine
        from apps.engine.pdf_passwords import clear_all_passwords

        clear_all_passwords()
        reset_engine()
        self._tmp.cleanup()
        if self.previous_localappdata is None:
            os.environ.pop("LOCALAPPDATA", None)
        else:
            os.environ["LOCALAPPDATA"] = self.previous_localappdata

    def _ingest(self, *paths: Path) -> dict:
        from apps.engine.dump import get_job, ingest_paths, start_job

        job = start_job(self.period["id"])
        ingest_paths(job["id"], [str(path) for path in paths])
        return get_job(job["id"]) or {}

    def _write_tally(self) -> Path:
        path = self.inbox / "Tally_Issues.xml"
        path.write_text(TALLY_WITH_ISSUES, encoding="utf-8")
        return path

    def test_pack_carries_the_preflight_node_and_warning_counts(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.pipeline import get_period_pack

        tally = self._write_tally()
        job = self._ingest(tally)
        self.assertNotEqual(job.get("status"), "failed", job)
        self.assertFalse(job.get("blocked"), job)

        pack = get_period_pack(self.period["id"])
        assert pack is not None
        self.assertFalse(pack.get("blocked"))
        node = pack["preflight"]
        self.assertTrue(node["ok"])
        counts = {
            issue["code"]: issue["count"]
            for issue in node["warnings"]
        }
        self.assertEqual(counts.get("hsn_missing"), 2)
        self.assertEqual(counts.get("missing_date"), 1)
        self.assertEqual(counts.get("duplicate_invoice"), 1)
        self.assertEqual(counts.get("invoice_math"), 1)
        self.assertEqual(node["blocking"], [])

        summary = json.loads((Path(pack["path"]) / "pack_summary.json").read_text(encoding="utf-8"))
        self.assertEqual(
            [issue["code"] for issue in summary["preflight"]["warnings"]],
            [issue["code"] for issue in node["warnings"]],
        )

        report_path = Path(pack["path"]) / "Preflight_Report.xlsx"
        self.assertTrue(report_path.is_file())
        sheet = load_workbook(report_path)["Preflight"]
        levels = [sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)]
        self.assertTrue(levels)
        self.assertEqual(set(levels), {"WARNING"})
        self.assertEqual(len(levels), len(node["warnings"]))

    def test_blocking_run_writes_nothing_and_keeps_the_existing_pack(self) -> None:
        from apps.engine.db import DataPack, get_session
        from apps.engine.dump import list_period_files, override_kind, reparse_period
        from apps.engine.pipeline import get_period_pack
        from apps.engine.tests.test_stage3_bank import write_hdfc

        bank = write_hdfc(self.inbox)
        self._ingest(bank)
        first_pack = get_period_pack(self.period["id"])
        assert first_pack is not None
        self.assertTrue(first_pack["exists"])
        pack_folder = Path(first_pack["path"])
        before = {
            path.name: (path.stat().st_mtime_ns, path.read_bytes())
            for path in pack_folder.iterdir()
        }
        summary_before = json.loads((pack_folder / "pack_summary.json").read_text(encoding="utf-8"))
        self.assertIn("preflight", summary_before)

        session = get_session()
        try:
            packs_before = session.query(DataPack).filter(DataPack.period_id == self.period["id"]).count()
        finally:
            session.close()

        # Drop the file back to unknown and re-run: nothing is left to pack.
        bank_file = next(row for row in list_period_files(self.period["id"]) if row["kind"] == "bank")
        override_kind(bank_file["id"], "unknown")
        job = reparse_period(self.period["id"])
        self.assertNotEqual(job.get("status"), "failed", job)
        self.assertTrue(job.get("blocked"), job)
        self.assertIn("no usable rows", (job.get("error_message") or "").lower())

        after = {
            path.name: (path.stat().st_mtime_ns, path.read_bytes())
            for path in pack_folder.iterdir()
        }
        self.assertEqual(after, before)

        # The preflight workbook in the folder still describes the old run.
        from openpyxl import load_workbook

        sheet = load_workbook(pack_folder / "Preflight_Report.xlsx")["Preflight"]
        self.assertEqual(sheet.max_row, 2)
        self.assertIn("nothing to report", str(sheet.cell(2, 3).value))

        pack = get_period_pack(self.period["id"])
        assert pack is not None
        self.assertTrue(pack["blocked"])
        self.assertEqual(
            [issue["code"] for issue in pack["preflight"]["blocking"]],
            ["no_usable_rows"],
        )

        session = get_session()
        try:
            packs_after = session.query(DataPack).filter(DataPack.period_id == self.period["id"]).count()
        finally:
            session.close()
        self.assertEqual(packs_after, packs_before)

        from apps.engine.preflight import load_report

        report = load_report(self.period["id"])
        assert report is not None
        self.assertTrue(report.blocked)

    def test_unwritable_destination_blocks_the_run(self) -> None:
        from unittest.mock import patch

        from apps.engine.pipeline import get_period_pack
        from apps.engine.tests.test_stage3_bank import write_hdfc

        blocker = Path(self._tmp.name) / "a-file"
        blocker.write_text("not a folder", encoding="utf-8")
        unreachable = blocker / "pack"

        with patch("apps.engine.pipeline.period_output_dir", return_value=unreachable):
            bank = write_hdfc(self.inbox)
            job = self._ingest(bank)

        self.assertTrue(job.get("blocked"), job)
        self.assertNotEqual(job.get("status"), "failed", job)
        self.assertIn("output folder cannot be written", (job.get("error_message") or "").lower())

        pack = get_period_pack(self.period["id"])
        assert pack is not None
        self.assertTrue(pack["blocked"])
        self.assertEqual(
            [issue["code"] for issue in pack["preflight"]["blocking"]],
            ["destination_unwritable"],
        )
        self.assertFalse(unreachable.exists())
        pack_root = Path(self._tmp.name) / "CAUnpacker" / "packs"
        written = [path for path in pack_root.rglob("*") if path.is_file()] if pack_root.exists() else []
        self.assertEqual(written, [])


if __name__ == "__main__":
    unittest.main()
