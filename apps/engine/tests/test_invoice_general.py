from __future__ import annotations

import os
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[3]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from apps.engine.tests import fixtures_invoice_layouts as layouts

FLPK = Path(__file__).resolve().parent / "fixtures" / "invoices" / "flpk.pdf"
OCR_ENGINES = {"rapidocr", "tesseract"}


class InvoiceFixtureTestCase(unittest.TestCase):
    def parse_pages(self, pages, *, font_size: float = 9.0) -> dict:
        from apps.engine.parsers.invoice import parse_invoice_file

        tmp = tempfile.TemporaryDirectory(prefix="caunpacker_invoice_fixture_")
        self.addCleanup(tmp.cleanup)
        path = layouts.write_pdf(Path(tmp.name) / "invoice.pdf", pages, font_size=font_size)
        return parse_invoice_file(path)

    def assert_checks_ok(self, row: dict) -> None:
        checks = row.get("checks") or []
        self.assertTrue(checks, row)
        self.assertEqual({check["status"] for check in checks}, {"ok"}, checks)

    def assert_no_math_break(self, row: dict) -> None:
        self.assertNotIn("invoice_math", row["flags"], row)
        for check in row.get("checks") or []:
            if check["check"] in {
                "taxable_plus_tax",
                "line_items_total",
                "line_items_taxable",
                "line_items_tax",
                "tax_components",
            }:
                self.assertEqual(check["status"], "ok", check)


@unittest.skipUnless(FLPK.is_file(), "flpk.pdf regression fixture is unavailable")
class FlpkRealInvoiceTests(InvoiceFixtureTestCase):
    def test_two_invoices_and_three_line_items(self) -> None:
        from apps.engine.parsers.invoice import parse_invoice_file

        parsed = parse_invoice_file(FLPK)
        self.assertFalse(parsed["unreadable"])
        self.assertNotEqual(parsed["engine"], "none")
        self.assertEqual(parsed["record_count"], 2)
        self.assertEqual(parsed["invoice_record_count"], 2)
        self.assertEqual(parsed["failed_records"], [])
        self.assertEqual(parsed["skipped_rows"], [])
        self.assertEqual(parsed["ignored_documents"], [])

        rows = {row["invoice_number"]: row for row in parsed["rows"]}
        self.assertEqual(set(rows), {"FAIH9X2500133941", "BFF9725004803384"})

        first = rows["FAIH9X2500133941"]
        self.assertEqual(first["invoice_date"], "2025-01-09")
        self.assertEqual(first["supplier_name"], "Health & Happiness Private Limited")
        self.assertEqual(first["supplier_gstin"], "29AADCH8449J1Z0")
        self.assertEqual(first["customer_name"], "Darshan V")
        self.assertTrue(first["customer_address"])
        self.assertEqual(first["currency"], "INR")
        self.assertEqual(first["hsn"], "84713010")
        self.assertAlmostEqual(first["taxable_value"], 30923.72, places=2)
        self.assertAlmostEqual(first["tax"], 5566.28, places=2)
        self.assertAlmostEqual(first["cgst"], 2783.14, places=2)
        self.assertAlmostEqual(first["sgst"], 2783.14, places=2)
        self.assertIsNone(first["igst"])
        self.assertAlmostEqual(first["invoice_value"], 36490.0, places=2)
        self.assertEqual(first["tax_rate"], 18.0)
        self.assert_checks_ok(first)

        second = rows["BFF9725004803384"]
        self.assertEqual(second["invoice_date"], "2025-01-09")
        self.assertEqual(second["supplier_gstin"], "29AADCH8449J1Z0")
        self.assertEqual(second["currency"], "INR")
        self.assertEqual(second["hsn"], "998599")
        self.assertAlmostEqual(second["taxable_value"], 24.58, places=2)
        self.assertAlmostEqual(second["tax"], 4.42, places=2)
        self.assertAlmostEqual(second["cgst"], 2.21, places=2)
        self.assertAlmostEqual(second["sgst"], 2.21, places=2)
        self.assertAlmostEqual(second["invoice_value"], 29.0, places=2)
        self.assertEqual(second["tax_rate"], 18.0)
        self.assert_checks_ok(second)

        items = parsed["line_items"]
        self.assertEqual(len(items), 3)
        counts = Counter(item["invoice_number"] for item in items)
        self.assertEqual(counts["FAIH9X2500133941"], 2)
        self.assertEqual(counts["BFF9725004803384"], 1)

        laptop = next(item for item in items if "ASUS Vivobook" in item["description"])
        self.assertEqual(laptop["invoice_number"], "FAIH9X2500133941")
        self.assertIn("E1504FA", laptop["description"])
        self.assertEqual(laptop["hsn"], "84713010")
        self.assertEqual(laptop["qty"], 1.0)
        self.assertAlmostEqual(laptop["gross"], 36490.0, places=2)
        self.assertAlmostEqual(laptop["discount"], 0.0, places=2)
        self.assertAlmostEqual(laptop["taxable"], 30923.72, places=2)
        self.assertAlmostEqual(laptop["cgst"], 2783.14, places=2)
        self.assertAlmostEqual(laptop["sgst"], 2783.14, places=2)
        self.assertAlmostEqual(laptop["tax"], 5566.28, places=2)
        self.assertAlmostEqual(laptop["amount"], 36490.0, places=2)

        shipping = next(item for item in items if "Shipping" in item["description"])
        self.assertEqual(shipping["invoice_number"], "FAIH9X2500133941")
        self.assertAlmostEqual(shipping["gross"], 70.0, places=2)
        self.assertAlmostEqual(shipping["discount"], -70.0, places=2)
        self.assertAlmostEqual(shipping["taxable"], 0.0, places=2)
        self.assertAlmostEqual(shipping["amount"], 0.0, places=2)

        protect = next(item for item in items if "Protect Promise" in item["description"])
        self.assertEqual(protect["invoice_number"], "BFF9725004803384")
        self.assertEqual(protect["hsn"], "998599")
        self.assertEqual(protect["qty"], 1.0)
        self.assertAlmostEqual(protect["taxable"], 24.58, places=2)
        self.assertAlmostEqual(protect["tax"], 4.42, places=2)
        self.assertAlmostEqual(protect["amount"], 29.0, places=2)

        self.assertFalse(
            any(
                item["description"].lower().startswith(("total", "grand total"))
                for item in items
            )
        )
        for item in items:
            self.assertEqual(item["currency"], "INR")
            self.assertEqual(item["source"], "flpk.pdf")


class MultipleInvoiceTests(InvoiceFixtureTestCase):
    def test_credit_note_is_separated_from_invoices(self) -> None:
        parsed = self.parse_pages(layouts.multi_invoice_pages())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(parsed["record_count"], 3)
        self.assertEqual(parsed["invoice_record_count"], 2)
        self.assertEqual(
            [row["invoice_number"] for row in parsed["rows"]],
            ["INV/AA/001", "INV/BB/002"],
        )
        self.assertEqual(len(parsed["ignored_documents"]), 1)
        ignored = parsed["ignored_documents"][0]
        self.assertEqual(ignored["document_type"], "credit_note")
        self.assertEqual(ignored["number"], "CR/001")
        self.assertIn("credit_note_ignored", parsed["warnings"])

        rows = {row["invoice_number"]: row for row in parsed["rows"]}
        first = rows["INV/AA/001"]
        self.assertEqual(first["invoice_date"], "2026-04-01")
        self.assertEqual(first["supplier_name"], "Alpha Traders")
        self.assertAlmostEqual(first["taxable_value"], 200.0, places=2)
        self.assertAlmostEqual(first["tax"], 36.0, places=2)
        self.assertAlmostEqual(first["invoice_value"], 236.0, places=2)
        self.assert_checks_ok(first)

        second = rows["INV/BB/002"]
        self.assertEqual(second["invoice_date"], "2026-04-02")
        self.assertEqual(second["supplier_name"], "Beta Supplies")
        self.assertAlmostEqual(second["invoice_value"], 500.0, places=2)
        self.assert_no_math_break(second)

        item_numbers = [item["invoice_number"] for item in parsed["line_items"]]
        self.assertEqual(item_numbers, ["INV/AA/001", "INV/BB/002"])
        self.assertNotIn("CR/001", item_numbers)

    def test_two_invoice_numbers_on_one_page_are_split(self) -> None:
        parsed = self.parse_pages(layouts.same_page_invoice_pages())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(parsed["record_count"], 2)
        self.assertEqual(
            [row["invoice_number"] for row in parsed["rows"]],
            ["SAME/001", "SAME/002"],
        )
        totals = {row["invoice_number"]: row["invoice_value"] for row in parsed["rows"]}
        self.assertEqual(totals, {"SAME/001": 118.0, "SAME/002": 236.0})
        by_invoice = {item["invoice_number"]: item for item in parsed["line_items"]}
        self.assertEqual(by_invoice["SAME/001"]["description"], "Item A")
        self.assertEqual(by_invoice["SAME/002"]["description"], "Item B")
        for row in parsed["rows"]:
            self.assert_no_math_break(row)

    def test_one_invoice_across_pages_with_repeated_header(self) -> None:
        parsed = self.parse_pages(layouts.multipage_invoice_pages())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(parsed["record_count"], 1)
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "MULTI/001")
        self.assertAlmostEqual(row["taxable_value"], 300.0, places=2)
        self.assertAlmostEqual(row["tax"], 54.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 354.0, places=2)
        self.assert_checks_ok(row)
        self.assertEqual(
            {item["description"] for item in parsed["line_items"]},
            {"Item A", "Item B"},
        )
        self.assertEqual(
            {item["invoice_number"] for item in parsed["line_items"]},
            {"MULTI/001"},
        )


class LayoutRegressionTests(InvoiceFixtureTestCase):
    def test_wrapped_rows_multiline_header_and_hsn_continuation(self) -> None:
        parsed = self.parse_pages(layouts.wrapped_table_page())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "WRAP/001")
        self.assertEqual(row["supplier_name"], "Wrap Industries Pvt Ltd")
        self.assertEqual(row["hsn"], "998314")
        self.assertAlmostEqual(row["taxable_value"], 1200.0, places=2)
        self.assertAlmostEqual(row["tax"], 216.0, places=2)
        self.assertAlmostEqual(row["cgst"], 108.0, places=2)
        self.assertAlmostEqual(row["sgst"], 108.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 1416.0, places=2)
        self.assert_checks_ok(row)

        items = parsed["line_items"]
        self.assertEqual(len(items), 2)
        first, second = items
        self.assertIn(
            "Widget with a very long description that wraps across three lines",
            first["description"],
        )
        self.assertEqual(first["hsn"], "998314")
        self.assertEqual(first["qty"], 1.0)
        self.assertAlmostEqual(first["gross"], 1180.0, places=2)
        self.assertAlmostEqual(first["discount"], 0.0, places=2)
        self.assertAlmostEqual(first["taxable"], 1000.0, places=2)
        self.assertAlmostEqual(first["cgst"], 90.0, places=2)
        self.assertAlmostEqual(first["sgst"], 90.0, places=2)
        self.assertAlmostEqual(first["tax"], 180.0, places=2)
        self.assertAlmostEqual(first["amount"], 1180.0, places=2)
        self.assertEqual(second["description"], "Service Fee")
        self.assertAlmostEqual(second["taxable"], 200.0, places=2)
        self.assertAlmostEqual(second["amount"], 236.0, places=2)

    def test_missing_optional_column_in_a_row(self) -> None:
        parsed = self.parse_pages(layouts.missing_column_page())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "MISS/001")
        self.assertAlmostEqual(row["taxable_value"], 245.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 245.0, places=2)
        self.assert_no_math_break(row)

        items = parsed["line_items"]
        self.assertEqual(len(items), 2)
        first, second = items
        self.assertEqual(first["description"], "Item A")
        self.assertEqual(first["qty"], 2.0)
        self.assertAlmostEqual(first["rate"], 100.0, places=2)
        self.assertIsNone(first["discount"])
        self.assertAlmostEqual(first["taxable"], 200.0, places=2)
        self.assertAlmostEqual(first["amount"], 200.0, places=2)
        self.assertEqual(second["description"], "Item B")
        self.assertAlmostEqual(second["discount"], 5.0, places=2)
        self.assertAlmostEqual(second["taxable"], 45.0, places=2)
        self.assertAlmostEqual(second["amount"], 45.0, places=2)

    def test_non_gst_usd_invoice_with_varied_labels(self) -> None:
        parsed = self.parse_pages(layouts.non_gst_page())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "NGP-2026-77")
        self.assertEqual(row["invoice_date"], "2026-04-03")
        self.assertEqual(row["supplier_name"], "Northwind Traders")
        self.assertIsNone(row["supplier_gstin"])
        self.assertIn("gstin_missing", row["flags"])
        self.assertEqual(row["currency"], "USD")
        self.assertIsNone(row["taxable_value"])
        self.assertIsNone(row["tax"])
        self.assertAlmostEqual(row["invoice_value"], 1250.0, places=2)
        self.assert_no_math_break(row)

        items = parsed["line_items"]
        self.assertEqual([item["description"] for item in items], ["Consulting", "Implementation"])
        self.assertEqual([item["amount"] for item in items], [500.0, 750.0])
        self.assertEqual(items[0]["qty"], 2.0)
        self.assertAlmostEqual(items[0]["rate"], 250.0, places=2)
        self.assertEqual(items[1]["qty"], 1.0)

    def test_header_row_is_not_used_as_supplier_name(self) -> None:
        parsed = self.parse_pages(layouts.header_name_trap_page())
        self.assertFalse(parsed["unreadable"])
        row = parsed["rows"][0]
        self.assertEqual(row["supplier_name"], "Acme Trading Private Limited")
        self.assertNotIn("Billing Address", row["supplier_name"])
        self.assertAlmostEqual(row["invoice_value"], 118.0, places=2)

    def test_indian_grouping_and_named_month_date(self) -> None:
        parsed = self.parse_pages(layouts.formatted_numbers_page())
        self.assertFalse(parsed["unreadable"])
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "FMT/001")
        self.assertEqual(row["invoice_date"], "2026-07-12")
        self.assertEqual(row["currency"], "INR")
        self.assertAlmostEqual(row["taxable_value"], 100000.0, places=2)
        self.assertAlmostEqual(row["tax"], 18000.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 118000.0, places=2)
        self.assert_checks_ok(row)
        item = parsed["line_items"][0]
        self.assertAlmostEqual(item["rate"], 100000.0, places=2)
        self.assertAlmostEqual(item["taxable"], 100000.0, places=2)
        self.assertAlmostEqual(item["amount"], 118000.0, places=2)

    def test_non_taxable_adjustment_is_reconciled(self) -> None:
        parsed = self.parse_pages(layouts.adjustment_page())
        self.assertFalse(parsed["unreadable"])
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "ADJ/001")
        self.assertAlmostEqual(row["adjustments"], -250.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 930.0, places=2)
        self.assert_checks_ok(row)
        checks = {check["check"]: check for check in row["checks"]}
        self.assertIn("non_taxable_adjustments", checks)
        self.assertAlmostEqual(checks["non_taxable_adjustments"]["expected"], -250.0, places=2)
        items = {item["description"]: item for item in parsed["line_items"]}
        self.assertEqual(set(items), {"Handset", "Exchange"})
        self.assertAlmostEqual(items["Exchange"]["taxable"], 0.0, places=2)
        self.assertAlmostEqual(items["Exchange"]["amount"], -250.0, places=2)
        self.assert_no_math_break(row)

    def test_line_item_starting_with_total_is_not_dropped(self) -> None:
        parsed = self.parse_pages(layouts.total_named_item_page())
        self.assertFalse(parsed["unreadable"])
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "TOT/001")
        self.assertAlmostEqual(row["invoice_value"], 590.0, places=2)
        self.assertEqual(len(parsed["line_items"]), 1)
        item = parsed["line_items"][0]
        self.assertEqual(item["description"], "Total Protection Plan")
        self.assertAlmostEqual(item["taxable"], 500.0, places=2)
        self.assertAlmostEqual(item["amount"], 590.0, places=2)
        self.assert_checks_ok(row)

    def test_inline_imei_metadata_is_not_part_of_description(self) -> None:
        parsed = self.parse_pages(layouts.metadata_continuation_page())
        self.assertFalse(parsed["unreadable"])
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "META/001")
        self.assertEqual(len(parsed["line_items"]), 1)
        item = parsed["line_items"][0]
        description = item["description"]
        self.assertEqual(description, "Portable Table YellowWood Black, 128 GB) Accessories")
        self.assertNotIn("IMEI", description)
        self.assertNotIn("SrNo", description)
        self.assertNotIn("Warranty", description)
        self.assertEqual(item["hsn"], "44219090")
        self.assertAlmostEqual(row["invoice_value"], 118.0, places=2)
        self.assert_no_math_break(row)

    def test_tax_summary_lines_before_total_are_not_items(self) -> None:
        parsed = self.parse_pages(layouts.tax_summary_before_total_page())
        self.assertFalse(parsed["unreadable"])
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "TAXSUM/001")
        self.assertEqual(parsed["skipped_rows"], [])
        self.assertNotIn("unparsed_rows", row["flags"])
        self.assertAlmostEqual(row["cgst"], 9.0, places=2)
        self.assertAlmostEqual(row["sgst"], 9.0, places=2)
        self.assertAlmostEqual(row["tax"], 18.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 118.0, places=2)
        self.assertEqual(len(parsed["line_items"]), 1)
        self.assertEqual(parsed["line_items"][0]["description"], "Item")
        self.assert_no_math_break(row)

    def test_ambiguous_row_is_reported_not_guessed(self) -> None:
        parsed = self.parse_pages(layouts.ambiguous_rows_page())
        self.assertFalse(parsed["unreadable"])
        self.assertEqual(len(parsed["rows"]), 1)
        row = parsed["rows"][0]
        self.assertEqual(row["invoice_number"], "AMB/001")
        self.assertIn("unparsed_rows", row["flags"])
        self.assertEqual(len(parsed["skipped_rows"]), 1)
        self.assertIn("Mystery", parsed["skipped_rows"][0]["text"])
        self.assertEqual(len(parsed["line_items"]), 1)
        self.assertEqual(parsed["line_items"][0]["description"], "Good Item")
        self.assertFalse(
            any("Mystery" in item["description"] for item in parsed["line_items"])
        )
        self.assertAlmostEqual(row["invoice_value"], 118.0, places=2)
        self.assert_no_math_break(row)


class UnsupportedLayoutTests(InvoiceFixtureTestCase):
    def test_prose_invoice_does_not_invent_rows(self) -> None:
        parsed = self.parse_pages(layouts.unsupported_page())
        self.assertTrue(parsed["unreadable"])
        self.assertEqual(parsed["rows"], [])
        self.assertEqual(parsed["line_items"], [])
        self.assertEqual(len(parsed["failed_records"]), 1)
        self.assertEqual(parsed["failed_records"][0]["reason"], "no line items")


class ScannedInvoiceTests(InvoiceFixtureTestCase):
    def test_scanned_pdf_is_read_with_local_ocr(self) -> None:
        from apps.engine.ocr import ocr_available
        from apps.engine.parsers.invoice import parse_invoice_file

        if not ocr_available():
            self.skipTest("local OCR is unavailable")

        tmp = tempfile.TemporaryDirectory(prefix="caunpacker_scanned_invoice_")
        self.addCleanup(tmp.cleanup)
        source = layouts.write_pdf(
            Path(tmp.name) / "scan-source.pdf",
            layouts.scanned_source_page(),
            font_size=12.0,
        )
        scan = layouts.write_scanned_pdf(source, Path(tmp.name) / "scan.pdf")
        parsed = parse_invoice_file(scan)
        if parsed["engine"] not in OCR_ENGINES:
            self.skipTest("no local OCR engine produced text")
        if not parsed["rows"]:
            self.skipTest("local OCR could not recover the synthetic scan")

        self.assertIn("ocr_source", parsed["warnings"])
        row = parsed["rows"][0]
        self.assertTrue(row["invoice_number"] and "001" in row["invoice_number"], row)
        self.assertEqual(row["invoice_date"], "2026-04-05")
        self.assertEqual(row["supplier_name"], "Scan Traders")
        self.assertAlmostEqual(row["taxable_value"], 1000.0, places=2)
        self.assertAlmostEqual(row["tax"], 180.0, places=2)
        self.assertAlmostEqual(row["igst"], 180.0, places=2)
        self.assertAlmostEqual(row["invoice_value"], 1180.0, places=2)
        self.assert_checks_ok(row)

        items = parsed["line_items"]
        self.assertEqual(len(items), 1)
        self.assertIn("Widget", items[0]["description"])
        self.assertAlmostEqual(items[0]["taxable"], 1000.0, places=2)
        self.assertAlmostEqual(items[0]["amount"], 1180.0, places=2)


@unittest.skipUnless(FLPK.is_file(), "flpk.pdf regression fixture is unavailable")
class FlpkPipelineTests(InvoiceFixtureTestCase):
    def setUp(self) -> None:
        self._tmp = tempfile.TemporaryDirectory(prefix="caunpacker_flpk_pipeline_")
        self._old_local = os.environ.get("LOCALAPPDATA")
        os.environ["LOCALAPPDATA"] = self._tmp.name
        from apps.engine.db import reset_engine

        reset_engine()

    def tearDown(self) -> None:
        from apps.engine.db import reset_engine

        reset_engine()
        if self._old_local is None:
            os.environ.pop("LOCALAPPDATA", None)
        else:
            os.environ["LOCALAPPDATA"] = self._old_local
        self._tmp.cleanup()

    def test_flpk_purchase_workbook_has_new_invoice_fields(self) -> None:
        from openpyxl import load_workbook

        from apps.engine.clients import create_client
        from apps.engine.dump import ingest_paths, start_job
        from apps.engine.firm import save_firm
        from apps.engine.periods import create_period
        from apps.engine.pipeline import get_period_pack

        save_firm("Test firm")
        client = create_client("Acme")
        period = create_period(client["id"], "Jan 2025")
        job = start_job(period["id"])
        source = Path(self._tmp.name) / "flpk.pdf"
        source.write_bytes(FLPK.read_bytes())
        result = ingest_paths(job["id"], [str(source)])
        self.assertEqual(result.get("status"), "done", result)

        pack = get_period_pack(period["id"])
        self.assertIsNotNone(pack)
        output = next(item for item in pack["outputs"] if item["key"] == "purchase")
        book = load_workbook(output["path"])
        invoices = book["Invoices"]
        items = book["Line items"]

        invoice_headers = [cell.value for cell in invoices[1]]
        for title in ("Customer", "Discount", "CGST", "SGST", "IGST", "Tax rate %", "Currency"):
            self.assertIn(title, invoice_headers)
        item_headers = [cell.value for cell in items[1]]
        for title in ("Gross", "Discount", "CGST", "SGST", "IGST", "Tax rate %", "Currency", "Flags"):
            self.assertIn(title, item_headers)

        self.assertEqual(invoices.max_row, 3)
        self.assertEqual(items.max_row, 4)

        invoice_rows = {
            row[invoice_headers.index("Invoice no")]: row
            for row in invoices.iter_rows(min_row=2, values_only=True)
        }
        first = invoice_rows["FAIH9X2500133941"]
        self.assertEqual(first[invoice_headers.index("Customer")], "Darshan V")
        self.assertAlmostEqual(first[invoice_headers.index("Taxable")], 30923.72, places=2)
        self.assertAlmostEqual(first[invoice_headers.index("CGST")], 2783.14, places=2)
        self.assertAlmostEqual(first[invoice_headers.index("SGST")], 2783.14, places=2)
        self.assertAlmostEqual(first[invoice_headers.index("Tax rate %")], 18.0, places=2)
        self.assertAlmostEqual(first[invoice_headers.index("Invoice value")], 36490.0, places=2)
        self.assertEqual(first[invoice_headers.index("Currency")], "INR")
        self.assertEqual(first[invoice_headers.index("HSN")], "84713010")

        item_rows = list(items.iter_rows(min_row=2, values_only=True))
        laptop = next(
            row
            for row in item_rows
            if "ASUS Vivobook" in str(row[item_headers.index("Description")])
        )
        self.assertEqual(laptop[item_headers.index("HSN")], "84713010")
        self.assertAlmostEqual(laptop[item_headers.index("Gross")], 36490.0, places=2)
        self.assertAlmostEqual(laptop[item_headers.index("Discount")], 0.0, places=2)
        self.assertAlmostEqual(laptop[item_headers.index("CGST")], 2783.14, places=2)
        self.assertAlmostEqual(laptop[item_headers.index("SGST")], 2783.14, places=2)
        self.assertAlmostEqual(laptop[item_headers.index("Tax rate %")], 18.0, places=2)
        self.assertAlmostEqual(laptop[item_headers.index("Amount")], 36490.0, places=2)
        self.assertEqual(laptop[item_headers.index("Currency")], "INR")

        cgst_column = invoice_headers.index("CGST") + 1
        gross_column = item_headers.index("Gross") + 1
        self.assertEqual(invoices.cell(2, cgst_column).number_format, "#,##0.00")
        self.assertEqual(items.cell(2, gross_column).number_format, "#,##0.00")


if __name__ == "__main__":
    unittest.main()
