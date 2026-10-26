"""Workbook writer for the per-client treatment review."""

from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="E8EDE9")
REVIEW_FILL = PatternFill("solid", fgColor="FFCDD2")
CHECK_FILL = PatternFill("solid", fgColor="FFE7B3")
INFO_FILL = PatternFill("solid", fgColor="EAF3FF")
MONEY_FORMAT = "#,##0.00"

FLAG_COLS = (
    ("Severity", "severity"),
    ("Area", "area"),
    ("Finding", "title"),
    ("What was seen", "detail"),
    ("Why it was raised", "basis"),
    ("Amounts", "observed_text"),
    ("Working limit", "threshold_text"),
    ("Law or rule", "law_ref"),
    ("What the CA should do", "action"),
)
HEAD_COLS = (
    ("Suggested head", "head"),
    ("Lines matched", "count"),
    ("Amount", "amount"),
    ("Matched words", "keywords_text"),
    ("Sample", "sample_text"),
)


def _money(value) -> str:
    if value is None:
        return ""
    return f"{float(value):,.2f}"


def _observed_text(observed: Mapping | None) -> str:
    if not observed:
        return ""
    parts: list[str] = []
    if observed.get("count") is not None:
        parts.append(f"{observed['count']} line(s)")
    for key, label in (("total", "total"), ("largest", "largest single"), ("turnover", "turnover")):
        if observed.get(key) is not None:
            parts.append(f"{label} ₹{_money(observed[key])}")
    if observed.get("sales_lines") is not None and observed.get("count") is None:
        parts.append(f"{observed['sales_lines']} sales line(s)")
    return " · ".join(parts)


def _threshold_text(threshold: Mapping | None) -> str:
    if not threshold:
        return ""
    amount = threshold.get("amount")
    if amount is None:
        return str(threshold.get("basis") or "")
    text = f"₹{_money(amount)}"
    if threshold.get("enhanced"):
        text += f" (enhanced ₹{_money(threshold['enhanced'])})"
    if threshold.get("basis"):
        text += f" {threshold['basis']}"
    if threshold.get("rate"):
        text += f" · rate {threshold['rate']}"
    if threshold.get("as_of"):
        text += f" · {threshold['as_of']}"
    return text.strip()


def _write_flags(sheet, flags: Sequence[Mapping]) -> None:
    sheet.title = "Treatment notes"
    sheet.append([title for title, _key in FLAG_COLS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for flag in flags:
        row = dict(flag)
        row["observed_text"] = _observed_text(flag.get("observed"))
        row["threshold_text"] = _threshold_text(flag.get("threshold"))
        sheet.append([row.get(key) or "" for _title, key in FLAG_COLS])
        excel_row = sheet.max_row
        fill = {"review": REVIEW_FILL, "check": CHECK_FILL, "info": INFO_FILL}.get(
            str(flag.get("severity"))
        )
        for col in range(1, len(FLAG_COLS) + 1):
            cell = sheet.cell(excel_row, col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill is not None:
                cell.fill = fill

    if not flags:
        sheet.append(["", "", "No treatment flags for this period."])

    widths = (11, 18, 34, 60, 32, 26, 30, 34, 42)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _write_heads(sheet, heads: Sequence[Mapping], unmapped: Mapping | None) -> None:
    sheet.title = "Expense heads"
    sheet.append([title for title, _key in HEAD_COLS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for head in heads:
        samples = head.get("samples") or []
        sample_text = " | ".join(
            f"{item.get('date') or ''} {item.get('party') or ''} {item.get('narration') or ''}".strip()
            for item in samples[:3]
        )
        sheet.append(
            [
                head.get("head") or "",
                head.get("count") or 0,
                head.get("amount"),
                ", ".join(head.get("keywords") or []),
                sample_text,
            ]
        )
        sheet.cell(sheet.max_row, 3).number_format = MONEY_FORMAT

    if unmapped and unmapped.get("count"):
        sheet.append(
            [
                "Not matched to a suggested head",
                unmapped.get("count"),
                unmapped.get("amount"),
                "",
                " | ".join(
                    f"{item.get('date') or ''} {item.get('party') or ''} {item.get('narration') or ''}".strip()
                    for item in (unmapped.get("samples") or [])[:3]
                ),
            ]
        )
        sheet.cell(sheet.max_row, 3).number_format = MONEY_FORMAT

    if sheet.max_row == 1:
        sheet.append(["No line items were available for head mapping."])

    for index, width in enumerate((34, 13, 16, 30, 70), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _write_profile(sheet, treatment: Mapping) -> None:
    sheet.title = "Profile"
    profile = treatment.get("profile") or {}
    rows = [
        ("Client profile", profile.get("summary") or "Not set"),
        ("Profession", profile.get("profession_label") or ""),
        ("Profession group", profile.get("profession_group") or ""),
        ("Entity type", profile.get("entity_type_label") or ""),
        (
            "GST registration",
            "Registered"
            if profile.get("gst_registered") is True
            else "Not registered"
            if profile.get("gst_registered") is False
            else "Not set",
        ),
        ("GSTIN", profile.get("gstin") or ""),
        ("Typical GST rate", profile.get("gst_rate") or ""),
        ("GST note", profile.get("gst_note") or ""),
        ("Watch for", profile.get("watch") or ""),
        ("Rules as of", treatment.get("as_of") or ""),
        ("How to read this sheet", treatment.get("disclaimer") or ""),
    ]
    details = profile.get("details") or {}
    for key in ("books", "banks", "broker", "notes"):
        value = details.get(key)
        if value:
            rows.append((key.capitalize(), ", ".join(value) if isinstance(value, list) else value))

    sheet.append(["Field", "Value"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for label, value in rows:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.cell(sheet.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 100
    sheet.freeze_panes = "A2"


def write_treatment_workbook(path: Path, treatment: Mapping) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    _write_flags(book.active, treatment.get("flags") or [])
    _write_heads(book.create_sheet("Expense heads"), treatment.get("expense_heads") or [], treatment.get("unmapped"))
    _write_profile(book.create_sheet("Profile"), treatment)
    book.save(path)
    return path
