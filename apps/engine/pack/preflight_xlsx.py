"""The Preflight_Report workbook: one row per finding, blocking rows in red.

Same look as every other pack - the shared header band, banded body, hairline
borders and the ``#,##0.00`` money format all come from ``pack.style``. The
only difference is that a blocking finding paints its row red so it cannot be
mistaken for a warning at a glance.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from apps.engine.pack.style import (
    RED_FILL,
    fill_row,
    print_setup,
    style_body,
    style_header,
    style_note,
)

PREFLIGHT_HEADERS = ["Level", "Code", "What is wrong", "Count", "Where to look"]
PREFLIGHT_WIDTHS = [12, 22, 62, 8, 54]
LEVEL_LABELS = {"blocking": "BLOCKING", "warning": "WARNING"}


def write_preflight_report(path: Path, report) -> Path:
    """Write *report* (a ``PreflightReport``) to *path* and return the path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Preflight"
    sheet.append(PREFLIGHT_HEADERS)
    style_header(sheet, len(PREFLIGHT_HEADERS))

    blocking_rows: list[int] = []
    for issue in report.issues:
        samples = "; ".join(sample.text() for sample in issue.samples)
        sheet.append(
            [
                LEVEL_LABELS.get(issue.level, issue.level.upper()),
                issue.code,
                issue.message,
                issue.count,
                samples,
            ]
        )
        if issue.level == "blocking":
            blocking_rows.append(sheet.max_row)

    if not report.issues:
        sheet.append(["", "", "Pre-flight found nothing to report.", 0, ""])
        style_note(sheet, f"A{sheet.max_row}")

    style_body(sheet, 2, max(sheet.max_row, 1), len(PREFLIGHT_HEADERS))
    for excel_row in blocking_rows:
        fill_row(sheet, excel_row, len(PREFLIGHT_HEADERS), RED_FILL)

    last_col = get_column_letter(len(PREFLIGHT_HEADERS))
    last_row = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{last_col}{last_row}"
    sheet.freeze_panes = "A2"
    for index, width in enumerate(PREFLIGHT_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    print_setup(sheet, len(PREFLIGHT_HEADERS), header_row=1)
    book.save(path)
    return path
