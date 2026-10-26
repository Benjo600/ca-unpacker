"""Workbook writer for the per-client file review sheet."""

from __future__ import annotations

from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="E8EDE9")
MISSING_FILL = PatternFill("solid", fgColor="FFCDD2")
CHECK_FILL = PatternFill("solid", fgColor="FFE7B3")
ATTACHED_FILL = PatternFill("solid", fgColor="EAF3FF")
OK_FILL = PatternFill("solid", fgColor="EAF3F0")

CHECKLIST_COLS = (
    ("Grade", "grade_text"),
    ("Document", "label"),
    ("Status", "status_text"),
    ("Evidence", "evidence"),
    ("Why it matters", "why"),
    ("What to do", "action"),
)

# The cross-checks sheet shares the treatment flag columns on purpose, so both
# sheets read the same way.
FLAG_COLS = (
    ("Severity", "severity"),
    ("Area", "area"),
    ("Finding", "title"),
    ("What was seen", "detail"),
    ("Why it was raised", "basis"),
    ("Amounts", "observed_text"),
    ("Working limit", "threshold_text"),
    ("Law or rule", "law_ref"),
    ("What the CA should do", "action"),
)

STATUS_TEXT = {
    "missing": "Not in the pack",
    "attached": "Attached, not read",
    "verified": "Read this period",
}

_GRADE_TEXT = {"essential": "Essential", "recommended": "Recommended"}


def _money(value) -> str:
    if value is None:
        return ""
    return f"{float(value):,.2f}"


def _observed_text(observed: Mapping | None) -> str:
    if not observed:
        return ""
    parts: list[str] = []
    label_map = (
        ("count", "line(s)"),
        ("total", "total"),
        ("difference", "difference"),
        ("purchase_total", "purchases"),
        ("portal_total", "portal"),
        ("books_total", "books"),
        ("bank_credits", "bank credits"),
        ("sales_total", "sales"),
        ("missing_days", "missing days"),
    )
    for key, label in label_map:
        value = observed.get(key)
        if value is None:
            continue
        if isinstance(value, float) or isinstance(value, int):
            if key in {"count", "missing_days", "gstin_count", "name_count"}:
                parts.append(f"{value} {label}")
            else:
                parts.append(f"{label} ₹{_money(value)}")
        else:
            parts.append(f"{label} {value}")
    extra_keys = [
        key
        for key in observed
        if key not in {k for k, _l in label_map} and observed.get(key) is not None
    ]
    for key in sorted(extra_keys)[:4]:
        parts.append(f"{key.replace('_', ' ')}: {observed[key]}")
    return " · ".join(parts)


def _threshold_text(threshold: Mapping | None) -> str:
    if not threshold:
        return ""
    parts = []
    amount = threshold.get("amount")
    if amount is not None:
        parts.append(f"₹{_money(amount)}")
    for key in ("enhanced", "tolerance", "basis", "rate", "gap_days", "as_of"):
        if threshold.get(key) is not None and key != "amount":
            parts.append(str(threshold[key]))
    return " · ".join(parts)


def _write_checklist(sheet, checklist: Sequence[Mapping]) -> None:
    sheet.title = "Checklist"
    sheet.append([title for title, _key in CHECKLIST_COLS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for item in checklist:
        row = dict(item)
        row["grade_text"] = _GRADE_TEXT.get(str(row.get("grade")), str(row.get("grade")))
        row["status_text"] = STATUS_TEXT.get(str(row.get("status")), str(row.get("status")))
        if row.get("escalated"):
            row["grade_text"] += " (upgraded)"
        sheet.append([row.get(key) or "" for _title, key in CHECKLIST_COLS])
        excel_row = sheet.max_row
        status = str(item.get("status"))
        grade = str(item.get("grade"))
        if status == "missing":
            fill = MISSING_FILL if grade == "essential" else CHECK_FILL
        elif status == "attached":
            fill = ATTACHED_FILL
        else:
            fill = OK_FILL
        for col in range(1, len(CHECKLIST_COLS) + 1):
            cell = sheet.cell(excel_row, col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = fill

    if not checklist:
        sheet.append(["", "", "No checklist items for this profile."])

    widths = (14, 34, 18, 52, 52, 48)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _write_checks(sheet, flags: Sequence[Mapping]) -> None:
    sheet.title = "Cross-checks"
    sheet.append([title for title, _key in FLAG_COLS])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for flag in flags:
        row = dict(flag)
        row["observed_text"] = _observed_text(flag.get("observed"))
        row["threshold_text"] = _threshold_text(flag.get("threshold"))
        sheet.append([row.get(key) or "" for _title, key in FLAG_COLS])
        excel_row = sheet.max_row
        fill = {
            "review": MISSING_FILL,
            "check": CHECK_FILL,
            "info": ATTACHED_FILL,
        }.get(str(flag.get("severity")))
        for col in range(1, len(FLAG_COLS) + 1):
            cell = sheet.cell(excel_row, col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill is not None:
                cell.fill = fill

    if not flags:
        sheet.append(["", "", "No cross-document problems found this period."])

    widths = (11, 22, 34, 60, 32, 26, 30, 40, 42)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _write_summary(sheet, review: Mapping) -> None:
    sheet.title = "Summary"
    readiness_text = {
        "ready": "Ready — the expected documents are in and the cross-checks came back clean.",
        "check": "Needs checking — some documents are attached but unread, or a cross-check needs an eye.",
        "gaps": "Gaps — something the file should contain is missing, or a cross-check needs work.",
    }.get(str(review.get("readiness")), str(review.get("readiness")))
    counts = review.get("counts") or {}
    profile = review.get("profile") or {}
    rows = [
        ("Readiness", readiness_text),
        ("Client profile", profile.get("summary") or ""),
        (
            "Missing essentials",
            counts.get("essential_missing", 0),
        ),
        ("Missing nice-to-haves", counts.get("recommended_missing", 0)),
        ("Attached but unread", counts.get("attached_not_read", 0)),
        ("Documents read this period", counts.get("verified", 0)),
        ("Cross-checks to review", counts.get("review", 0)),
        ("Cross-checks to verify", counts.get("check", 0)),
        ("Cross-check notes", counts.get("info", 0)),
        ("Rules as of", review.get("as_of") or ""),
        ("How to read this sheet", review.get("disclaimer") or ""),
    ]
    sheet.append(["Field", "Value"])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
    for label, value in rows:
        sheet.append([label, value])
        sheet.cell(sheet.max_row, 1).font = Font(bold=True)
        sheet.cell(sheet.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 100
    sheet.freeze_panes = "A2"


def write_file_review_workbook(path: Path, review: Mapping) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    _write_summary(book.active, review)
    _write_checklist(book.create_sheet("Checklist"), review.get("checklist") or [])
    _write_checks(book.create_sheet("Cross-checks"), review.get("checks") or [])
    book.save(path)
    return path
