from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from apps.engine.parsers.bank.money import MONEY_TOKEN_RE
from apps.engine.parsers.sap.signatures import (
    CORE_FIELDS,
    FIELD_ALIASES,
    FINANCIAL_FIELDS,
    field_for_header,
    fields_for_headers,
    normalize_header,
)

_HEADER_SCAN_ROWS = 80
_SAP_DATE_RE = re.compile(
    r"\b("
    r"\d{4}-\d{1,2}-\d{1,2}"
    r"|"
    r"\d{1,2}[/\-.]\d{1,2}[/\-.]\d{2,4}"
    r"|"
    r"\d{1,2}\s*[A-Za-z]{3,9}\.?\s*\d{2,4}"
    r")\b"
)
_TOTAL_RE = re.compile(r"^\s*(?:grand\s+)?totals?\b|^\s*sum\b", re.I)
_CURRENCY_RE = re.compile(r"\b[A-Z]{3}\b")
_INDICATOR_TOKEN_RE = re.compile(
    r"^(?:S|H|D|C|DR|CR|DEBIT|CREDIT|SOLL|HABEN)$",
    re.I,
)
_ID_TOKEN_RE = re.compile(r"^[A-Z0-9][A-Z0-9/.\-]{2,}$", re.I)
_TEXT_FIELDS = (
    "party_name",
    "account_name",
    "description",
    "document_type",
    "reference",
    "assignment",
    "currency",
)
_MONEY_FIELDS = (
    "debit",
    "credit",
    "amount",
    "balance",
    "opening_balance",
    "closing_balance",
    "aging_current",
    "aging_1_30",
    "aging_31_60",
    "aging_61_90",
    "aging_91_120",
    "aging_over_120",
)
_ID_FIELDS = (
    "company_code",
    "document_number",
    "gl_account",
    "vendor_number",
    "customer_number",
    "account_number",
    "party_code",
    "posting_key",
    "fiscal_year",
    "period",
)


@dataclass
class TableRecord:
    values: dict[str, Any]
    raw_values: dict[str, str]
    raw_cells: list[str]
    page: int = 1
    bbox: str | None = None
    is_total: bool = False

    @property
    def raw_text(self) -> str:
        return " | ".join(cell for cell in self.raw_cells if cell).strip()


@dataclass
class ParsedTable:
    headers: list[str]
    field_headers: dict[str, str]
    records: list[TableRecord]
    context: str
    sheet: str | None = None
    kind: str = "table"


@dataclass
class _Token:
    text: str
    x: float
    right: float
    word_index: int = 0


@dataclass
class _HeaderColumn:
    field: str
    label: str
    x: float
    right: float

    @property
    def center(self) -> float:
        return (self.x + self.right) / 2.0


@dataclass
class _WordRow:
    page: int
    words: list
    text: str
    bbox: str | None


_ALIAS_TOKENS: list[tuple[tuple[str, ...], str]] = []
for _field, _aliases in FIELD_ALIASES.items():
    for _alias in _aliases:
        _tokens = tuple(normalize_header(_alias).split())
        if not _tokens:
            continue
        if (_tokens, _field) not in _ALIAS_TOKENS:
            _ALIAS_TOKENS.append((_tokens, _field))
_ALIAS_TOKENS.sort(key=lambda item: (-len(item[0]), -len(" ".join(item[0]))))


def read_tables(path: Path, filename: str | None = None) -> dict:
    """Read all report tables from a CSV, XLSX or PDF without raising."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return {
            "tables": [],
            "engine": "none",
            "error": None,
            "reason": "Legacy .xls is not supported — export the SAP report as .xlsx or .csv.",
        }
    try:
        if suffix == ".csv":
            tables = _tables_from_csv(path)
            return {"tables": tables, "engine": "csv", "error": None}
        if suffix == ".xlsx":
            tables = _tables_from_xlsx(path)
            return {"tables": tables, "engine": "xlsx", "error": None}
        if suffix == ".pdf":
            return _tables_from_pdf(path, filename)
        return {
            "tables": [],
            "engine": "none",
            "error": None,
            "reason": "SAP reports must be CSV, XLSX, PDF or a supported spreadsheet export.",
        }
    except Exception as exc:  # keep file-level failures visible, never crash the dump
        return {
            "tables": [],
            "engine": "none",
            "error": str(exc) or "Could not read the SAP report.",
        }


def read_table_samples(path: Path, max_rows: int = 40) -> list[tuple[str | None, list[list[Any]]]]:
    """Return sheet-sized header samples for classification.

    This is deliberately separate from the full parser so the router never
    loads a large workbook into memory.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [(None, _read_csv_matrix(path, max_rows=max_rows))]
    if suffix == ".xlsx":
        return _read_xlsx_samples(path, max_rows=max_rows)
    return []


def find_header_row(
    matrix: list[list[Any]],
) -> tuple[int, dict[int, str], set[str], int] | None:
    best: tuple[int, dict[int, str], set[str], int] | None = None
    for index, row in enumerate(matrix[:_HEADER_SCAN_ROWS]):
        cells = [_cell_text(cell) for cell in row]
        fields = fields_for_headers(cells)
        score = _table_header_score(fields)
        if score <= 0:
            continue
        field_map = {
            cell_index: field
            for cell_index, cell in enumerate(cells)
            if (field := field_for_header(cell)) is not None
        }
        if best is None or score > best[3]:
            best = (index, field_map, fields, score)
    return best


def _table_header_score(fields: set[str]) -> int:
    if len(fields) < 2 or not (fields & FINANCIAL_FIELDS):
        return 0
    return (
        len(fields) * 3
        + len(fields & CORE_FIELDS) * 2
        + len(fields & FINANCIAL_FIELDS) * 2
    )


def _tables_from_csv(path: Path) -> list[ParsedTable]:
    matrix = _read_csv_matrix(path)
    table = _table_from_matrix(matrix, sheet=None, kind="csv")
    return [table] if table else []


def _tables_from_xlsx(path: Path) -> list[ParsedTable]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[ParsedTable] = []
    try:
        for sheet in workbook.worksheets:
            matrix = [list(row) for row in sheet.iter_rows(values_only=True)]
            table = _table_from_matrix(matrix, sheet=sheet.title, kind="xlsx")
            if table is not None:
                tables.append(table)
    finally:
        workbook.close()
    return tables


def _read_csv_matrix(path: Path, max_rows: int | None = None) -> list[list[str]]:
    text = _read_text(path)
    if not text.strip():
        return []
    delimiter = _choose_delimiter(text)
    rows: list[list[str]] = []
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    for index, row in enumerate(reader):
        if max_rows is not None and index >= max_rows:
            break
        rows.append([str(cell) for cell in row])
    return rows


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except (UnicodeError, UnicodeDecodeError):
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _choose_delimiter(text: str) -> str:
    sample_lines = text.splitlines()[:8]
    counts = {
        delimiter: sum(line.count(delimiter) for line in sample_lines)
        for delimiter in (",", ";", "\t", "|")
    }
    delimiter = max(counts, key=counts.get)
    return delimiter if counts[delimiter] else ","


def _read_xlsx_samples(
    path: Path, max_rows: int = 40
) -> list[tuple[str | None, list[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return []
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    samples: list[tuple[str | None, list[list[Any]]]] = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= max_rows:
                    break
                rows.append(list(row))
            samples.append((sheet.title, rows))
    finally:
        workbook.close()
    return samples


def _table_from_matrix(
    matrix: list[list[Any]], *, sheet: str | None, kind: str
) -> ParsedTable | None:
    found = find_header_row(matrix)
    if found is None:
        return None
    header_index, field_map, _fields, _score = found
    header_cells = [_cell_text(cell) for cell in matrix[header_index]]
    field_headers = {
        field: header_cells[index] or field
        for index, field in field_map.items()
        if index < len(header_cells)
    }
    context = _matrix_context(matrix[:header_index])
    records: list[TableRecord] = []
    for row in matrix[header_index + 1 :]:
        cells = [_cell_text(cell) for cell in row]
        if not any(cell for cell in cells):
            continue
        if _is_repeated_header_cells(cells):
            continue
        values: dict[str, Any] = {}
        raw_values: dict[str, str] = {}
        for index, field in field_map.items():
            value = row[index] if index < len(row) else None
            if _is_empty(value):
                continue
            values[field] = value
            raw_values[field_headers.get(field, field)] = _cell_text(value)
        if not values:
            continue
        records.append(
            TableRecord(
                values=values,
                raw_values=raw_values,
                raw_cells=cells,
                page=1,
                bbox=None,
                is_total=_row_is_total(cells),
            )
        )
    return ParsedTable(
        headers=header_cells,
        field_headers=field_headers,
        records=records,
        context=context,
        sheet=sheet,
        kind=kind,
    )


def _matrix_context(rows: list[list[Any]]) -> str:
    lines: list[str] = []
    for row in rows:
        line = " ".join(_cell_text(cell) for cell in row if _cell_text(cell))
        if line:
            lines.append(line)
    return "\n".join(lines)[:4000]


def _tables_from_pdf(path: Path, filename: str | None) -> dict:
    from apps.engine.pdf_extract import extract_pdf

    extracted = extract_pdf(path)
    if extracted.pdf_type == "encrypted":
        return {
            "tables": [],
            "engine": extracted.engine,
            "error": None,
            "pdf_type": "encrypted",
            "reason": "SAP PDF is password-protected. Unlock it and parse again.",
        }
    word_rows = _word_rows(extracted.words)
    if word_rows:
        table = _table_from_word_rows(word_rows, extracted)
        if table is not None and _word_table_has_transactions(table):
            return {
                "tables": [table],
                "engine": extracted.engine,
                "error": None,
                "pdf_type": extracted.pdf_type,
            }
    table = _table_from_text_lines(extracted)
    if table is not None:
        return {
            "tables": [table],
            "engine": extracted.engine,
            "error": None,
            "pdf_type": extracted.pdf_type,
        }
    return {
        "tables": [],
        "engine": extracted.engine,
        "error": None,
        "pdf_type": extracted.pdf_type,
        "reason": "No supported SAP table could be read from this PDF layout.",
    }


def _word_rows(words: list) -> list[_WordRow]:
    if not words:
        return []
    by_page: dict[int, list] = {}
    for word in words:
        text = (getattr(word, "text", "") or "").strip()
        if not text:
            continue
        by_page.setdefault(int(getattr(word, "page", 1)), []).append(word)
    rows: list[_WordRow] = []
    for page, page_words in by_page.items():
        ordered = sorted(page_words, key=lambda item: (-float(item.y), float(item.x)))
        bands: list[list] = []
        for word in ordered:
            placed = False
            for band in bands:
                tolerance = max(2.0, max(float(word.height), float(band[0].height)) * 0.55)
                if abs(float(word.y) - float(band[0].y)) <= tolerance:
                    band.append(word)
                    placed = True
                    break
            if not placed:
                bands.append([word])
        for band in bands:
            band.sort(key=lambda item: float(item.x))
            text = " ".join((getattr(word, "text", "") or "").strip() for word in band)
            if not text.strip():
                continue
            xs = [float(word.x) for word in band]
            ys = [float(word.y) for word in band]
            rights = [float(word.x) + float(word.width) for word in band]
            tops = [float(word.y) + float(word.height) for word in band]
            bbox = (
                f"{min(xs):.1f},{min(ys):.1f},"
                f"{max(rights) - min(xs):.1f},{max(tops) - min(ys):.1f}"
            )
            rows.append(_WordRow(page=page, words=band, text=text, bbox=bbox))
    rows.sort(key=lambda row: (row.page, -_bbox_top(row.bbox)))
    return rows


def _bbox_top(bbox: str | None) -> float:
    if not bbox:
        return 0.0
    try:
        _x, y, _width, height = (float(part) for part in bbox.split(","))
        return y + height
    except (TypeError, ValueError):
        return 0.0


def _table_from_word_rows(word_rows: list[_WordRow], extracted) -> ParsedTable | None:
    best_index = None
    best_columns: list[_HeaderColumn] = []
    best_score = 0.0
    for index, row in enumerate(word_rows):
        columns = _segment_word_header(row)
        score = _header_columns_score(columns)
        if score > best_score:
            best_score = score
            best_index = index
            best_columns = columns
    if best_index is None or best_score <= 0 or len(best_columns) < 2:
        return None

    columns = sorted(best_columns, key=lambda column: column.x)
    field_headers: dict[str, str] = {}
    for column in columns:
        field_headers.setdefault(column.field, column.label)
    context = _pdf_context(word_rows[:best_index], extracted)
    records: list[TableRecord] = []
    for row in word_rows[best_index + 1 :]:
        repeated = _segment_word_header(row)
        if len(repeated) >= 2 and _header_columns_score(repeated) >= best_score * 0.9:
            continue
        values: dict[str, Any] = {}
        raw_values: dict[str, str] = {}
        cells_by_field: dict[str, list[str]] = {}
        for word in row.words:
            text = (getattr(word, "text", "") or "").strip()
            if not text:
                continue
            column = _column_for_word(word, columns)
            if column is None:
                continue
            cells_by_field.setdefault(column.field, []).append(text)
        for field, parts in cells_by_field.items():
            value = " ".join(parts).strip()
            if value:
                values[field] = value
                raw_values[field_headers.get(field, field)] = value
        if not values:
            continue
        raw_cells = [" ".join(cells_by_field.get(column.field, [])) for column in columns]
        records.append(
            TableRecord(
                values=values,
                raw_values=raw_values,
                raw_cells=raw_cells,
                page=row.page,
                bbox=row.bbox,
                is_total=_row_is_total(raw_cells),
            )
        )
    return ParsedTable(
        headers=[column.label for column in columns],
        field_headers=field_headers,
        records=records,
        context=context,
        sheet=None,
        kind="pdf",
    )


def _segment_word_header(row: _WordRow) -> list[_HeaderColumn]:
    tokens: list[_Token] = []
    for word_index, word in enumerate(row.words):
        text = normalize_header(getattr(word, "text", ""))
        if not text:
            continue
        x = float(word.x)
        right = float(word.x) + float(word.width)
        for piece in text.split():
            tokens.append(_Token(text=piece, x=x, right=right, word_index=word_index))
    return _segment_header_tokens(tokens)


def _segment_header_tokens(tokens: list[_Token]) -> list[_HeaderColumn]:
    columns: list[_HeaderColumn] = []
    seen: set[str] = set()
    index = 0
    while index < len(tokens):
        matched = None
        for alias_tokens, field in _ALIAS_TOKENS:
            length = len(alias_tokens)
            if index + length > len(tokens):
                continue
            if tuple(token.text for token in tokens[index : index + length]) != alias_tokens:
                continue
            if length > 1 and any(
                tokens[position + 1].x - tokens[position].right > 20.0
                for position in range(index, index + length - 1)
            ):
                # "Debit" and "Credit" are separate columns, not one
                # "Debit/Credit" indicator column.
                continue
            if alias_tokens == ("debit", "credit") and len(
                {token.word_index for token in tokens[index : index + length]}
            ) > 1:
                # Two separate header words next to each other are amount
                # columns.  Only a single word such as "Debit/Credit" is an
                # indicator column.
                continue
            if field in seen:
                continue
            matched = (field, length)
            break
        if matched is None:
            index += 1
            continue
        field, length = matched
        span = tokens[index : index + length]
        columns.append(
            _HeaderColumn(
                field=field,
                label=" ".join(token.text for token in span),
                x=min(token.x for token in span),
                right=max(token.right for token in span),
            )
        )
        seen.add(field)
        index += length
    return columns


def _word_table_has_transactions(table: ParsedTable) -> bool:
    """A word table is only useful when rows keep several mapped columns.

    Some PDF engines return a whole text line as one "word".  Such a line is
    assigned entirely to the first header column, so it must be parsed by the
    line-based fallback instead of being accepted as a table row.
    """
    return any(
        len(record.values) >= 2 and bool(set(record.values) & set(_MONEY_FIELDS))
        for record in table.records
    )


def _header_columns_score(columns: list[_HeaderColumn]) -> float:
    fields = {column.field for column in columns}
    if len(fields) < 2:
        return 0.0
    financial = len(fields & FINANCIAL_FIELDS)
    if financial < 1:
        return 0.0
    return len(fields) * 3.0 + financial * 2.0


def _column_for_word(word, columns: list[_HeaderColumn]) -> _HeaderColumn | None:
    if not columns:
        return None
    center = float(word.x) + float(word.width) / 2.0
    if center <= columns[0].center:
        return columns[0]
    for index, column in enumerate(columns[:-1]):
        boundary = (column.center + columns[index + 1].center) / 2.0
        if center < boundary:
            return column
    return columns[-1]


def _pdf_context(rows: list[_WordRow], extracted) -> str:
    lines = [row.text for row in rows if row.text.strip()]
    if not lines:
        lines = [
            page.text.strip()
            for page in (getattr(extracted, "pages", None) or [])[:2]
            if (page.text or "").strip()
        ]
    return "\n".join(lines)[:4000]


def _table_from_text_lines(extracted) -> ParsedTable | None:
    lines = [
        line
        for line in (getattr(extracted, "lines", None) or [])
        if (getattr(line, "text", "") or "").strip()
    ]
    if not lines:
        return None
    best_index = None
    best_columns: list[_HeaderColumn] = []
    best_score = 0.0
    for index, line in enumerate(lines):
        columns = _segment_text_header(line.text)
        score = _header_columns_score(columns)
        if score > best_score:
            best_score = score
            best_index = index
            best_columns = columns
    if best_index is None or best_score <= 0 or len(best_columns) < 2:
        return None

    fields = [column.field for column in best_columns]
    field_headers = {column.field: column.label for column in best_columns}
    records: list[TableRecord] = []
    for line in lines[best_index + 1 :]:
        text = (line.text or "").strip()
        if not text:
            continue
        repeated = _segment_text_header(text)
        if len(repeated) >= 2 and _header_columns_score(repeated) >= best_score * 0.9:
            continue
        values = _parse_text_data_line(text, fields)
        if not values:
            continue
        records.append(
            TableRecord(
                values=values,
                raw_values={field_headers[field]: str(value) for field, value in values.items()},
                raw_cells=[text],
                page=int(getattr(line, "page", 1) or 1),
                bbox=_line_bbox(line),
                is_total=bool(_TOTAL_RE.match(text)),
            )
        )
    return ParsedTable(
        headers=[column.label for column in best_columns],
        field_headers=field_headers,
        records=records,
        context="\n".join((line.text or "") for line in lines[:best_index])[:4000],
        sheet=None,
        kind="pdf-text",
    )


def _segment_text_header(text: str) -> list[_HeaderColumn]:
    tokens: list[_Token] = []
    position = 0.0
    word_index = 0
    for raw in str(text or "").split():
        normalized = normalize_header(raw)
        if not normalized:
            continue
        for piece in normalized.split():
            tokens.append(
                _Token(
                    text=piece,
                    x=position,
                    right=position + 10.0,
                    word_index=word_index,
                )
            )
            position += 10.0
        word_index += 1
    return _segment_header_tokens(tokens)


def _parse_text_data_line(text: str, fields: list[str]) -> dict[str, Any]:
    remaining = text
    values: dict[str, Any] = {}

    date_matches = list(_SAP_DATE_RE.finditer(remaining))
    date_values = [match.group(1) for match in date_matches]
    remaining = _remove_spans(remaining, [match.span() for match in date_matches])

    money_matches = list(MONEY_TOKEN_RE.finditer(remaining))
    money_values = [match.group(0) for match in money_matches]
    remaining = _remove_spans(remaining, [match.span() for match in money_matches])

    currency_values = _CURRENCY_RE.findall(remaining)
    if "currency" in fields:
        for currency in currency_values:
            remaining = re.sub(rf"\b{re.escape(currency)}\b", " ", remaining, count=1)

    tokens = [token for token in remaining.split() if token]
    if money_values and "debit_credit_indicator" in fields:
        candidates = [token for token in tokens if _INDICATOR_TOKEN_RE.match(token)]
        # More than one short Dr/Cr-looking token is ambiguous (for example a
        # description containing "C"), so do not guess which one is the flag.
        if len(candidates) == 1:
            values["debit_credit_indicator"] = candidates[0]
            tokens.remove(candidates[0])
    numeric_values: list[str] = []
    text_values: list[str] = []
    for token in tokens:
        if _ID_TOKEN_RE.match(token) and any(char.isdigit() for char in token):
            numeric_values.append(token)
        else:
            text_values.append(token)

    date_fields = [
        field
        for field in fields
        if field in {"posting_date", "document_date", "entry_date", "due_date", "baseline_date", "clearing_date"}
    ]
    for field, value in zip(date_fields, date_values):
        values[field] = value

    money_fields = [field for field in fields if field in _MONEY_FIELDS]
    if money_values:
        if len(money_values) == len(money_fields):
            for field, value in zip(money_fields, money_values):
                values[field] = value
        elif len(money_fields) == 1:
            values[money_fields[0]] = money_values[0]
        elif "amount" in money_fields and len(money_values) == 1:
            values["amount"] = money_values[0]
        # A different count means the text layout omitted or merged columns.
        # Leaving the amounts unassigned is safer than guessing debit vs credit.

    id_fields = [field for field in fields if field in _ID_FIELDS]
    for field, value in zip(id_fields, numeric_values):
        values[field] = value

    if currency_values and "currency" in fields:
        values["currency"] = currency_values[0]

    text_fields = [field for field in fields if field in _TEXT_FIELDS]
    if text_values and text_fields:
        # Text columns cannot be separated safely without word coordinates.
        # Preserve the complete text against the first textual column instead
        # of inventing a split between party, description and reference.
        values[text_fields[0]] = " ".join(text_values)

    return values


def _remove_spans(text: str, spans: list[tuple[int, int]]) -> str:
    if not spans:
        return text
    chars = list(text)
    for start, end in spans:
        for index in range(max(0, start), min(len(chars), end)):
            chars[index] = " "
    return "".join(chars)


def _line_bbox(line) -> str | None:
    try:
        x = float(line.x)
        y = float(line.y)
        width = float(line.width)
        height = float(line.height)
    except (TypeError, ValueError):
        return None
    if width <= 0 or height <= 0:
        return None
    return f"{x:.1f},{y:.1f},{width:.1f},{height:.1f}"


def _is_repeated_header_cells(cells: list[str]) -> bool:
    fields = fields_for_headers(cells)
    return len(fields) >= 3 and bool(fields & FINANCIAL_FIELDS)


def _row_is_total(cells: list[str]) -> bool:
    first = next((cell for cell in cells if cell), "")
    return bool(_TOTAL_RE.match(first))


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())
